"""ÜYBS referansları ve rapor Excel çıktıları."""

import datetime
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from timetable import refs, reports
from timetable.models import Card, Lesson, LessonRef, Period, School, SchoolClass, Subject, Teacher, Timetable


class RefTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name="MYO", days_per_week=5)
        for i in range(6):
            Period.objects.create(school=self.school, index=i, name=f"{i + 1}")
        self.cls = SchoolClass.objects.create(school=self.school, name="1.KISIM", short="1SVN")
        teacher = Teacher.objects.create(school=self.school, first_name="Ali", last_name="Kaya")
        self.lesson = Lesson.objects.create(
            school=self.school, subject=Subject.objects.create(school=self.school, name="Askerî Tarih", short="ASK102"),
            periods_per_week=2, card_length=2,
        )
        self.lesson.classes.set([self.cls])
        self.lesson.teachers.set([teacher])
        self.timetable = Timetable.objects.create(school=self.school, name="t", status=Timetable.OK)
        self.card = Card.objects.create(timetable=self.timetable, lesson=self.lesson, length=2, day=0, period=0)
        School.objects.filter(pk=self.school.pk).update(active_timetable=self.timetable)

    def test_refs_follow_the_lesson_when_moved(self):
        self.assertEqual(refs.auto_number(self.school, self.timetable, "1001"), 2)
        rows = refs.report_rows(self.timetable)
        self.assertEqual([(r["day"], r["period"], r["ref"]) for r in rows], [("Pazartesi", 1, "1001"), ("Pazartesi", 2, "1002")])
        # Başka bir hafta ders 3. ve 4. saate alındı
        self.card.period = 2
        self.card.save()
        rows = refs.report_rows(self.timetable, week_start=datetime.date(2026, 11, 2))
        self.assertEqual([(r["period"], r["ref"]) for r in rows], [(3, "1001"), (4, "1002")])
        self.assertEqual(rows[0]["date"], datetime.date(2026, 11, 2))
        self.assertEqual(rows[0]["class"], "1SVN")
        self.assertEqual(rows[0]["code"], "ASK102")

    def test_split_cards_keep_chronological_order(self):
        self.lesson.periods_per_week = 3
        self.lesson.card_length = 1
        self.lesson.save()
        self.card.length = 1
        self.card.day, self.card.period = 2, 0  # Çarşamba
        self.card.save()
        Card.objects.create(timetable=self.timetable, lesson=self.lesson, length=1, day=0, period=4)  # Pazartesi
        Card.objects.create(timetable=self.timetable, lesson=self.lesson, length=1)  # yerleşmemiş
        refs.import_text(self.school, "1SVN\tASK102\tA1, A2, A3")
        rows = refs.report_rows(self.timetable)
        self.assertEqual([(r["day"], r["ref"]) for r in rows], [("Pazartesi", "A1"), ("Çarşamba", "A2"), ("yerleşmedi", "A3")])

    def test_import_text_and_errors(self):
        written, errors = refs.import_text(self.school, "1SVN;ASK102;2;X2\n1svn;askerî tarih;1;X1\n9Z;ASK102;1;Y\n1SVN;ASK102;5;Z")
        self.assertEqual(written, 2)
        self.assertEqual(len(errors), 2)
        self.assertEqual(dict(LessonRef.objects.values_list("ordinal", "ref")), {1: "X1", 2: "X2"})

    def test_auto_number_keeps_existing_unless_overwrite(self):
        refs.save_cell(self.lesson, self.cls, 1, "777")
        refs.auto_number(self.school, self.timetable, "A0100", step=10)
        self.assertEqual(dict(LessonRef.objects.values_list("ordinal", "ref")), {1: "777", 2: "A0100"})
        refs.auto_number(self.school, self.timetable, "A0100", step=10, overwrite=True)
        self.assertEqual(dict(LessonRef.objects.values_list("ordinal", "ref")), {1: "A0100", 2: "A0110"})

    def test_page_report_and_excel(self):
        user = User.objects.create_superuser("uybs", password="parola-123456")
        self.client.force_login(user)
        url = reverse("timetable:lesson_refs", args=[self.school.pk])
        page = self.client.get(url)
        self.assertContains(page, "Pzt 1")
        self.assertContains(page, "0 / 2 saat")
        response = self.client.post(url, {"action": "save", f"ref-{self.lesson.pk}-{self.cls.pk}-1": "1001", f"ref-{self.lesson.pk}-{self.cls.pk}-2": "1002"})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(LessonRef.objects.count(), 2)
        report = self.client.get(reverse("timetable:report_render", args=[self.timetable.pk]) + "?tur=uybs&hafta=2026-11-04")
        self.assertContains(report, "02.11.2026")  # Çarşamba verildi, Pazartesi'ye çekildi
        self.assertContains(report, "<b>1001</b>")
        xlsx = self.client.get(reverse("timetable:report_excel", args=[self.timetable.pk]) + "?tur=uybs&hafta=2026-11-02")
        ws = load_workbook(BytesIO(xlsx.content)).active  # Raporlar'daki Excel de ÜYBS şablonu biçiminde
        self.assertEqual((ws.title, ws["A1"].value, ws["C1"].value), ("UYBS_Ders_Aktarim", "ÜYBS REF", "DERS KODU"))
        self.assertEqual([ws["A2"].value, ws["A3"].value], [1001, 1002])

    def test_every_report_type_exports_to_excel(self):
        user = User.objects.create_superuser("excel2", password="parola-123456")
        self.client.force_login(user)
        for kind in reports.TYPES:
            response = self.client.get(reverse("timetable:report_excel", args=[self.timetable.pk]) + f"?tur={kind}")
            self.assertEqual(response.status_code, 200, kind)
            self.assertTrue(load_workbook(BytesIO(response.content)).worksheets, kind)


class RefInTimetablePrintTests(RefTests):
    def test_ref_shown_in_class_timetable(self):
        refs.auto_number(self.school, self.timetable, "1001")
        user = User.objects.create_superuser("refprint", password="parola-123456")
        self.client.force_login(user)
        page = self.client.get(reverse("timetable:report_render", args=[self.timetable.pk]) + "?tur=bireysel&varlik=sinif&icerik=ref&icerik_set=1")
        self.assertContains(page, "ref 1001")
        self.assertContains(page, "ref 1002")

"""Veritabanı, servis, API ve görünüm katmanının uçtan uca testleri."""

import datetime
import json
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from timetable import asc_xml, editing, substitutes
from timetable.demo import create_demo_school
from timetable.models import Absence, Availability, Card, CardRelationship, Lesson, Substitution, Timetable
from timetable.services import check_school, generate_timetable, repair_timetable, start_generation, verify_timetable


class DemoSchoolTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.school = create_demo_school()

    def test_demo_data_passes_checks(self):
        self.assertFalse([i for i in check_school(self.school) if i.level == "error"])

    def test_generate_places_all_cards_without_violations(self):
        timetable = generate_timetable(self.school, "test", time_limit=30)
        self.assertEqual(timetable.status, Timetable.OK, timetable.report)
        expected = sum(len(lesson.card_lengths()) for lesson in self.school.lessons.all())
        self.assertEqual(timetable.cards.filter(day__isnull=False).count(), expected)
        violations, stats = verify_timetable(timetable)
        self.assertEqual(violations, [])
        self.assertEqual(stats["unplaced"], 0)

    def test_physics_teacher_has_no_friday_lessons(self):
        timetable = generate_timetable(self.school, "test", time_limit=30)
        physics = timetable.cards.filter(lesson__subject__short="FİZ")
        self.assertTrue(physics.exists())
        self.assertFalse(physics.filter(day=4).exists())

    def test_part_time_teacher_comes_two_days(self):
        timetable = generate_timetable(self.school, "test", time_limit=30)
        days = set(timetable.cards.filter(lesson__teachers__short="MPO").values_list("day", flat=True))
        self.assertLessEqual(len(days), 2)

    def test_locked_cards_are_kept_when_regenerating(self):
        first = generate_timetable(self.school, "ilk", time_limit=30)
        card = first.cards.filter(lesson__subject__short="TDE").first()
        card.locked = True
        card.save()
        second = generate_timetable(self.school, "ikinci", time_limit=30, base=first)
        self.assertEqual(second.status, Timetable.OK)
        kept = second.cards.get(lesson=card.lesson, day=card.day, period=card.period)
        self.assertTrue(kept.locked)

    def test_lessons_added_after_base_are_included(self):
        first = generate_timetable(self.school, "ilk", time_limit=30)
        lesson = Lesson.objects.create(
            school=self.school, subject=self.school.subjects.get(short="MÜZ"), periods_per_week=1
        )
        lesson.classes.set([self.school.classes.get(name="10-A")])
        lesson.teachers.set([self.school.teachers.get(short="GAK")])
        second = generate_timetable(self.school, "ikinci", time_limit=30, base=first)
        self.assertEqual(second.status, Timetable.OK)
        self.assertEqual(second.cards.count(), first.cards.count() + 1)
        self.assertTrue(second.cards.get(lesson=lesson).is_placed)

    def test_manual_clash_is_detected_by_verification(self):
        timetable = generate_timetable(self.school, "test", time_limit=30)
        a, b = timetable.cards.filter(lesson__subject__short="TDE", lesson__classes__name="9-A")[:2]
        b.day, b.period = a.day, a.period
        b.save()
        violations, _ = verify_timetable(timetable)
        self.assertIn("atom_clash", [v.code for v in violations])

    def test_overloaded_teacher_fails_strict_but_relaxed_gives_draft(self):
        lesson = Lesson.objects.filter(school=self.school, subject__short="MAT").first()
        teacher = lesson.teachers.first()
        Availability.objects.bulk_create(
            Availability(school=self.school, teacher=teacher, day=d, period=p) for d in range(4) for p in range(8)
        )
        strict = generate_timetable(self.school, "katı", time_limit=5)
        self.assertEqual(strict.status, Timetable.FAILED)
        self.assertIn("teacher_overload", [i["code"] for i in strict.report["issues"]])
        self.assertFalse(Card.objects.filter(timetable=strict).exists())

        relaxed = generate_timetable(self.school, "gevşek", time_limit=15, mode=Timetable.RELAXED)
        self.assertEqual(relaxed.status, Timetable.PARTIAL)
        self.assertTrue(relaxed.cards.filter(day__isnull=True).exists())
        # Gevşek mod kural kırabilir ama kaynak çakışması (aynı öğretmen/sınıf/derslik aynı anda) asla olmaz.
        self.assertEqual([v for v in verify_timetable(relaxed)[0] if v.code.endswith("_clash")], [])

    def test_infeasible_strict_generation_is_explained(self):
        teacher = self.school.teachers.get(short="EŞA")  # Cuma gelmiyor, haftada 9 saat
        teacher.max_periods_per_day = 2
        teacher.save()
        timetable = generate_timetable(self.school, "açıklama", time_limit=20)
        self.assertEqual(timetable.status, Timetable.FAILED)
        conflict = timetable.report["conflict"]
        self.assertEqual(conflict["status"], "core")
        labels = sorted(rule["label"] for rule in conflict["rules"])
        self.assertEqual(labels, ["Elif Şahin (öğretmen): müsaitlik (yasak saatler)", "Elif Şahin: günde en fazla 2 ders"])
        self.assertTrue(all(rule["url"].endswith(f"/ogretmenler/{teacher.pk}/") for rule in conflict["rules"]))
        user = User.objects.create_user("aciklama", password="parola-123456")
        self.client.force_login(user)
        self.assertContains(self.client.get(reverse("timetable:timetable_detail", args=[timetable.pk])), "Neden çözülemiyor?")

    def test_teacher_days_conflict_is_reported_before_solving(self):
        teacher = self.school.teachers.get(short="AYI")
        teacher.max_days_per_week = 4
        teacher.save()
        self.assertIn("teacher_days", [i.code for i in check_school(self.school)])

    def test_hard_relationship_is_respected(self):
        rel = CardRelationship.objects.create(
            school=self.school, name="Fizik/Kimya farklı gün", condition=CardRelationship.DIFFERENT_DAYS,
            apply_to=CardRelationship.EACH_CLASS, importance=CardRelationship.STRICT,
        )
        rel.subjects.set(self.school.subjects.filter(short__in=["FİZ", "KİM"]))
        timetable = generate_timetable(self.school, "ilişki", time_limit=30)
        self.assertEqual(timetable.status, Timetable.OK)
        for cls in self.school.classes.all():
            days = list(
                timetable.cards.filter(lesson__classes=cls, lesson__subject__short__in=["FİZ", "KİM"]).values_list("day", flat=True)
            )
            self.assertEqual(len(days), len(set(days)), cls)


class EditingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.school = create_demo_school()
        cls.base = generate_timetable(cls.school, "düzenleme", time_limit=30)

    def setUp(self):
        self.timetable = Timetable.objects.get(pk=self.base.pk)

    def test_options_cover_whole_grid(self):
        card = self.timetable.cards.filter(length=1).first()
        options = editing.Board(self.timetable).options(card.id)
        self.assertEqual(len(options), 5 * 8)
        statuses = {o["status"] for o in options}
        self.assertTrue(statuses & {editing.FREE, editing.WARN})
        current = next(o for o in options if (o["day"], o["period"]) == (card.day, card.period))
        self.assertNotEqual(current["status"], editing.BLOCKED)

    def test_blocked_cell_for_teacher_time_off(self):
        card = self.timetable.cards.filter(lesson__subject__short="FİZ").first()
        options = editing.Board(self.timetable).options(card.id)
        friday = [o for o in options if o["day"] == 4]
        self.assertTrue(all(o["status"] == editing.BLOCKED for o in friday))
        result = editing.move_card(card, 4, 0)
        self.assertFalse(result["ok"])

    def test_move_into_conflict_displaces_cards(self):
        cards = list(self.timetable.cards.filter(lesson__classes__name="9-A", length=1, lesson__subject__short="TDE"))
        a, b = cards[0], cards[1]
        # a'yı, b'nin bulunduğu hücreye taşı: aynı sınıf → b bekleyenlere düşmeli.
        board = editing.Board(self.timetable)
        info = board.evaluate(a.id, b.day, b.period)
        if info["status"] == editing.BLOCKED:
            self.skipTest("Bu yerleşimde hedef hücre başka bir nedenle kapalı.")
        result = editing.move_card(a, b.day, b.period)
        self.assertTrue(result["ok"])
        self.assertIn(b.id, result["displaced"])
        b.refresh_from_db()
        self.assertFalse(b.is_placed)

    def test_unplace_and_repair(self):
        cards = list(self.timetable.cards.filter(lesson__classes__name="10-B")[:3])
        for card in cards:
            editing.unplace_card(card)
        self.assertEqual(self.timetable.cards.filter(day__isnull=True).count(), 3)
        result = repair_timetable(self.timetable, time_limit=15)
        self.assertTrue(result["ok"])
        self.assertEqual(self.timetable.cards.filter(day__isnull=True).count(), 0)
        self.assertEqual([v for v in verify_timetable(self.timetable)[0] if v.level == "hard"], [])

    def test_locked_card_cannot_be_unplaced(self):
        card = self.timetable.cards.first()
        card.locked = True
        card.save()
        self.assertFalse(editing.unplace_card(card)["ok"])


class ExportImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.school = create_demo_school()
        cls.timetable = generate_timetable(cls.school, "aktarım", time_limit=30)

    def test_asc_xml_round_trip(self):
        data = asc_xml.export_school(self.school, self.timetable)
        self.assertIn(b"<timetable", data)
        school, timetable = asc_xml.import_school(data, name="kopya")
        self.assertEqual(school.days_per_week, 5)
        self.assertEqual(school.periods.count(), 8)
        self.assertEqual(school.teachers.count(), self.school.teachers.count())
        self.assertEqual(school.lessons.count(), self.school.lessons.count())
        self.assertEqual(Lesson.objects.filter(school=school, groups__isnull=False).distinct().count(), 4)
        self.assertEqual(timetable.status, Timetable.OK)
        self.assertEqual(timetable.cards.count(), self.timetable.cards.count())
        self.assertEqual(sum(c.length for c in timetable.cards.all()), sum(c.length for c in self.timetable.cards.all()))
        self.assertEqual(verify_timetable(timetable)[0], [])

    def test_invalid_xml_is_rejected(self):
        with self.assertRaises(asc_xml.AscImportError):
            asc_xml.import_school(b"<html></html>")
        with self.assertRaises(asc_xml.AscImportError):
            asc_xml.import_school(b"not xml")

    def test_excel_export(self):
        from timetable import excel

        workbook = load_workbook(BytesIO(excel.export(self.timetable)))
        self.assertEqual(workbook.sheetnames, ["Sınıflar", "Öğretmenler", "Derslikler", "Kartlar"])
        self.assertEqual(workbook["Sınıflar"]["A4"].value, "9-A")
        self.assertEqual(workbook["Kartlar"].max_row, self.timetable.cards.count() + 1)


class SubstitutionTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.school = create_demo_school()
        cls.timetable = generate_timetable(cls.school, "vekâlet", time_limit=30)
        cls.school.active_timetable = cls.timetable
        cls.school.save()

    def test_plan_suggests_free_teachers(self):
        date = datetime.date(2026, 11, 2)  # Pazartesi
        teacher = self.timetable.cards.filter(day=0).first().lesson.teachers.first()
        absence = Absence.objects.create(school=self.school, date=date, teacher=teacher)
        plan = substitutes.day_plan(self.school, self.timetable, date)
        slots = plan[0]["slots"]
        expected = sum(c.length for c in self.timetable.cards.filter(day=0, lesson__teachers=teacher))
        self.assertEqual(len(slots), expected)
        for slot in slots:
            for candidate in slot["candidates"]:
                busy = self.timetable.cards.filter(day=0, lesson__teachers=candidate["teacher"])
                self.assertFalse(any(c.period <= slot["period"] < c.period + c.length for c in busy))
                self.assertNotEqual(candidate["teacher"], teacher)

        slot = slots[0]
        chosen = slot["candidates"][0]["teacher"]
        saved = substitutes.save_assignments(absence, self.timetable, {f"slot-{slot['key']}": str(chosen.pk)})
        self.assertEqual(saved, 1)
        self.assertEqual(Substitution.objects.get(absence=absence).substitute, chosen)
        # Kaydedilen vekil kendi atamasında aday olarak kalır; aynı saat için başka bir devamsızlıkta önerilmez.
        plan = substitutes.day_plan(self.school, self.timetable, date)
        self.assertIn(chosen, [c["teacher"] for c in plan[0]["slots"][0]["candidates"]])
        other_teacher = next(
            c.lesson.teachers.first()
            for c in self.timetable.cards.filter(day=0).exclude(lesson__teachers=teacher)
            if c.covers(0, slot["period"])
        )
        Absence.objects.create(school=self.school, date=date, teacher=other_teacher)
        plan = {item["absence"].teacher_id: item for item in substitutes.day_plan(self.school, self.timetable, date)}
        same_period = [s for s in plan[other_teacher.id]["slots"] if s["period"] == slot["period"]]
        self.assertTrue(same_period)
        self.assertNotIn(chosen, [c["teacher"] for c in same_period[0]["candidates"]])


@override_settings(TIMETABLE_RUN_IN_BACKGROUND=False)
class ViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.school = create_demo_school()
        cls.timetable = generate_timetable(cls.school, "görünüm", time_limit=30)
        cls.user = User.objects.create_user("mudur", password="parola-123456")

    def setUp(self):
        self.client.force_login(self.user)

    def test_login_is_required(self):
        self.client.logout()
        response = self.client.get(reverse("timetable:school_list"))
        self.assertEqual(response.status_code, 302)
        self.assertIn("/admin/login/", response["Location"])

    def test_pages_render(self):
        self.assertContains(self.client.get(reverse("timetable:school_list")), self.school.name)
        self.assertContains(self.client.get(reverse("timetable:school_detail", args=[self.school.pk])), "Yeni program üret")
        self.assertContains(self.client.get(reverse("timetable:timetable_detail", args=[self.timetable.pk])), "tt-board")
        self.assertContains(self.client.get(reverse("timetable:print", args=[self.timetable.pk]), {"tur": "ogretmen"}), "Pazartesi")
        self.assertContains(self.client.get(reverse("timetable:substitutions", args=[self.school.pk]), {"tarih": "2026-11-02"}), "Devamsızlık ekle")

    def test_grid_api(self):
        for kind, count in (("sinif", 4), ("ogretmen", 15), ("derslik", 9)):
            data = self.client.get(reverse("timetable:api_grid", args=[self.timetable.pk]), {"tur": kind}).json()
            self.assertEqual(len(data["rows"]), count)
            self.assertEqual(len(data["days"]), 5)
            self.assertEqual(len(data["periods"]), 8)
            self.assertTrue(all(c["rows"] for c in data["cards"]))
        data = self.client.get(reverse("timetable:api_grid", args=[self.timetable.pk])).json()
        self.assertEqual(data["stats"]["unplaced"], 0)
        self.assertEqual(data["violations"], [])

    def test_card_api_move_unplace_lock(self):
        # Program sıkı olduğundan her kartın boş bir hücresi olmayabilir; boş hücresi olan bir kart seç.
        for card in self.timetable.cards.filter(length=1):
            options = self.client.get(reverse("timetable:api_card_options", args=[card.pk])).json()["options"]
            target = next((o for o in options if o["status"] == "free" and (o["day"], o["period"]) != (card.day, card.period)), None)
            if target:
                break
        self.assertIsNotNone(target)
        response = self.client.post(
            reverse("timetable:api_card_move", args=[card.pk]),
            json.dumps({"day": target["day"], "period": target["period"]}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        card.refresh_from_db()
        self.assertEqual((card.day, card.period), (target["day"], target["period"]))

        self.assertEqual(self.client.post(reverse("timetable:api_card_lock", args=[card.pk])).json()["locked"], True)
        response = self.client.post(reverse("timetable:api_card_unplace", args=[card.pk]))
        self.assertEqual(response.status_code, 400)
        self.client.post(reverse("timetable:api_card_lock", args=[card.pk]))
        self.assertEqual(self.client.post(reverse("timetable:api_card_unplace", args=[card.pk])).status_code, 200)
        card.refresh_from_db()
        self.assertFalse(card.is_placed)

        response = self.client.post(reverse("timetable:api_repair", args=[self.timetable.pk]))
        self.assertTrue(response.json()["ok"])
        card.refresh_from_db()
        self.assertTrue(card.is_placed)

    def test_generate_view_runs_job_and_redirects(self):
        response = self.client.post(
            reverse("timetable:generate", args=[self.school.pk]), {"name": "web", "time_limit": "20", "mode": "strict"}
        )
        new = Timetable.objects.get(name="web")
        self.assertRedirects(response, reverse("timetable:timetable_detail", args=[new.pk]))
        self.assertEqual(new.status, Timetable.OK)
        status = self.client.get(reverse("timetable:api_status", args=[new.pk])).json()
        self.assertEqual(status["status"], "ok")

    def test_publish_and_exports(self):
        self.client.post(reverse("timetable:publish", args=[self.timetable.pk]))
        self.school.refresh_from_db()
        self.assertEqual(self.school.active_timetable, self.timetable)
        response = self.client.get(reverse("timetable:export_excel", args=[self.timetable.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        response = self.client.get(reverse("timetable:export_asc", args=[self.timetable.pk]))
        self.assertIn(b"<lessons", response.content)

    def test_import_view(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile("okulum.xml", asc_xml.export_school(self.school, self.timetable), content_type="text/xml")
        response = self.client.post(reverse("timetable:import_asc"), {"file": upload})
        imported = Timetable.objects.exclude(school=self.school).get()
        self.assertRedirects(response, reverse("timetable:timetable_detail", args=[imported.pk]))

    def test_substitution_flow(self):
        url = reverse("timetable:substitutions", args=[self.school.pk])
        teacher = self.timetable.cards.filter(day=0).first().lesson.teachers.first()
        self.client.post(url, {"action": "add", "teacher": teacher.pk, "tarih": "2026-11-02", "reason": "Raporlu"})
        self.assertTrue(Absence.objects.filter(teacher=teacher, reason="Raporlu").exists())
        response = self.client.get(url, {"tarih": "2026-11-02"})
        self.assertContains(response, "Raporlu")
        self.assertContains(response, "Önerilenleri seç")


class BackgroundJobTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.school = create_demo_school()

    @override_settings(TIMETABLE_RUN_IN_BACKGROUND=False)
    def test_start_generation_records_progress(self):
        timetable = start_generation(self.school, "iş", time_limit=20)
        timetable.refresh_from_db()
        self.assertEqual(timetable.status, Timetable.OK)
        self.assertIsNotNone(timetable.finished_at)
        self.assertGreaterEqual(timetable.report["solver"]["solutions"], 1)

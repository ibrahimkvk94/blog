"""Sınav programı: derslerden sınav üretme, ortak sınav birleştirme, çözücünün kesin kuralları, yetersiz kapasite
iletisi, görünümler, Excel çıktısı ve erişim tabloları."""

import datetime
from collections import defaultdict
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse

from timetable import access, exams, exams_solver
from timetable.models import (
    Availability, Classroom, Exam, ExamPeriod, ExamPlacement, ExamRoom, Lesson, Membership, Period, School, SchoolClass,
    Subject, Teacher,
)

MONDAY = datetime.date(2026, 11, 2)


def build_exam_school(classes=3, room_caps=(60, 60, 60), teachers=6):
    school = School.objects.create(name="Sınav okulu", days_per_week=5)
    for i in range(8):
        Period.objects.create(school=school, index=i, name=f"{i + 1}. saat", start_time=datetime.time(8 + i, 30),
                              end_time=datetime.time(9 + i, 20))
    tchs = [Teacher.objects.create(school=school, first_name=f"Öğretmen{i}", last_name="Test") for i in range(teachers)]
    for i, cap in enumerate(room_caps):
        Classroom.objects.create(school=school, name=f"Salon {i + 1}", capacity=cap)
    cls = [SchoolClass.objects.create(school=school, name=f"{i + 1}.KISIM", student_count=30) for i in range(classes)]
    common = Subject.objects.create(school=school, name="Matematik")
    for i, c in enumerate(cls):  # her sınıfa ayrı kısım: ortak sınavda birleşmeli
        lesson = Lesson.objects.create(school=school, subject=common, periods_per_week=2)
        lesson.classes.set([c])
        lesson.teachers.set([tchs[i % 2]])
    for j, name in enumerate(["Fizik", "Kimya"]):
        subject = Subject.objects.create(school=school, name=name)
        for i, c in enumerate(cls):
            lesson = Lesson.objects.create(school=school, subject=subject, periods_per_week=2)
            lesson.classes.set([c])
            lesson.teachers.set([tchs[2 + j]])
    Subject.objects.create(school=school, name="Boş ders")  # dersi olmayan ders adı: sınav üretilmez
    period = ExamPeriod.objects.create(school=school, name="2026 Güz Vize", start_date=MONDAY,
                                       end_date=MONDAY + datetime.timedelta(days=4), sessions=["09:30", "11:00", "13:30"],
                                       time_limit=10)
    return school, period


class ExamGenerationTests(TestCase):
    def setUp(self):
        self.school, self.period = build_exam_school()

    def test_common_exam_merges_sections(self):
        created = exams.generate_from_lessons(self.period)
        self.assertEqual(created, 3)
        math = Exam.objects.get(period=self.period, subject__name="Matematik")
        self.assertEqual(math.classes.count(), 3)
        self.assertEqual(math.lessons.count(), 3)
        self.assertEqual(math.student_count, 90)
        self.assertEqual(set(math.teachers.values_list("first_name", flat=True)), {"Öğretmen0", "Öğretmen1"})
        self.assertFalse(Exam.objects.filter(subject__name="Boş ders").exists())

    def test_separate_and_split_and_merge(self):
        physics = Subject.objects.get(school=self.school, name="Fizik")
        created = exams.generate_from_lessons(self.period, separate=[physics.pk])
        self.assertEqual(created, 2 + 3)
        self.assertEqual(Exam.objects.filter(period=self.period, subject=physics).count(), 3)
        ids = list(Exam.objects.filter(period=self.period, subject=physics).values_list("pk", flat=True))
        merged = exams.merge(self.period, ids)
        self.assertEqual(merged.classes.count(), 3)
        self.assertEqual(Exam.objects.filter(period=self.period, subject=physics).count(), 1)
        self.assertEqual(exams.split(self.period, merged.pk), 3)
        parts = Exam.objects.filter(period=self.period, subject=physics)
        self.assertEqual(sorted(e.student_count for e in parts), [30, 30, 30])
        # yalnızca eksik ders adlarını ekle: yeni sınav oluşmaz
        self.assertEqual(exams.generate_from_lessons(self.period, replace=False), 0)

    def test_student_records_override_class_size(self):
        from timetable.models import Student

        cls = SchoolClass.objects.get(school=self.school, name="1.KISIM")
        for i in range(5):
            Student.objects.create(school=self.school, first_name=f"Ö{i}", school_class=cls)
        exams.generate_from_lessons(self.period)
        self.assertEqual(Exam.objects.get(period=self.period, subject__name="Fizik").student_count, 65)


class ExamSolverTests(TestCase):
    def setUp(self):
        self.school, self.period = build_exam_school(classes=3)
        exams.generate_from_lessons(self.period)

    def assert_hard_constraints(self):
        period = self.period
        by_class_slot, by_room_slot, by_teacher_slot = set(), set(), set()
        per_day = defaultdict(int)
        rows = exams.placements(period)
        self.assertEqual(len(rows), period.exams.filter(excluded=False).count())
        for pl in rows:
            e = pl.exam
            cap = sum(exams.exam_capacity(r.classroom, period) for r in pl.rooms.all())
            self.assertGreaterEqual(cap, e.student_count, e)
            self.assertEqual(sum(r.seats for r in pl.rooms.all()), e.student_count)
            for c in e.classes.all():
                per_day[(c.pk, pl.date)] += 1
            for k in range(e.duration):
                for c in e.classes.all():
                    key = (c.pk, pl.date, pl.session + k)
                    self.assertNotIn(key, by_class_slot, "sınıfın aynı oturumda iki sınavı olamaz")
                    by_class_slot.add(key)
                for r in pl.rooms.all():
                    key = (r.classroom_id, pl.date, pl.session + k)
                    self.assertNotIn(key, by_room_slot, "salonda aynı anda tek sınav")
                    by_room_slot.add(key)
                    self.assertEqual(r.proctors.count(), period.proctors_per_room)
                    for t in r.proctors.all():
                        key = (t.pk, pl.date, pl.session + k)
                        self.assertNotIn(key, by_teacher_slot, "gözetmen aynı anda iki salonda olamaz")
                        by_teacher_slot.add(key)
        self.assertLessEqual(max(per_day.values()), period.max_per_day)
        self.assertEqual([i for i in exams.validate(period) if i["level"] == "error"], [])

    @override_settings(TIMETABLE_RUN_IN_BACKGROUND=False)
    def test_solver_respects_hard_constraints(self):
        self.period.max_per_day = 1
        self.period.save()
        # Matematik 90 öğrenci: %50 sınav kapasitesiyle (30+30+30) üç salona bölünmeli
        self.assertTrue(exams.start(self.period))
        self.period.refresh_from_db()
        self.assertEqual(self.period.status, ExamPeriod.OK, self.period.report)
        math = ExamPlacement.objects.get(exam__subject__name="Matematik")
        self.assertEqual(math.rooms.count(), 3)
        # ortak sınavda her sınıf ayrı salona düşer (30 öğrenci, 30 kişilik salonlar)
        mapping = exams.class_rooms(math, exams.class_sizes(self.school))
        self.assertEqual(sorted(len(v) for v in mapping.values()), [1, 1, 1])
        self.assertEqual(len({v[0].pk for v in mapping.values()}), 3)
        self.assert_hard_constraints()
        days = {pl.date for pl in exams.placements(self.period)}
        self.assertEqual(len(days), 3, "günde en fazla 1 sınav: üç ayrı gün")

    @override_settings(TIMETABLE_RUN_IN_BACKGROUND=False)
    def test_unavailable_proctor_and_fixed_exam(self):
        teacher = Teacher.objects.get(school=self.school, first_name="Öğretmen5")
        # Öğretmen5 Pazartesi hiçbir saatte gelemez; Öğretmen4 dönem boyunca muaf
        for p in range(8):
            Availability.objects.create(school=self.school, teacher=teacher, day=0, period=p)
        self.period.proctor_exempt.add(Teacher.objects.get(school=self.school, first_name="Öğretmen4"))
        chem = Exam.objects.get(period=self.period, subject__name="Kimya")
        chem.fixed_date, chem.fixed_session = MONDAY, 1
        chem.save()
        exams.start(self.period)
        self.period.refresh_from_db()
        self.assertIn(self.period.status, (ExamPeriod.OK, ExamPeriod.PARTIAL), self.period.report)
        pl = ExamPlacement.objects.get(exam=chem)
        self.assertEqual((pl.date, pl.session), (MONDAY, 1))
        monday = ExamRoom.objects.filter(placement__date=MONDAY)
        self.assertFalse(monday.filter(proctors=teacher).exists())
        self.assertFalse(ExamRoom.objects.filter(proctors__first_name="Öğretmen4").exists())
        self.assert_hard_constraints()

    def test_too_little_capacity_is_reported(self):
        Classroom.objects.filter(school=self.school).update(capacity=20)  # %50 → 10 + 10 + 10 = 30 < 90
        problem, _dates = exams.build_problem(self.period)
        result = exams_solver.solve(problem)
        self.assertEqual(result.status, "failed")
        self.assertTrue(any("kapasitesi yetersiz" in e and "Matematik" in e for e in result.errors), result.errors)

    def test_class_with_too_many_exams_is_reported(self):
        self.period.end_date = MONDAY
        self.period.max_per_day = 2
        self.period.save()
        problem, _dates = exams.build_problem(self.period)
        result = exams_solver.solve(problem)
        self.assertEqual(result.status, "failed")
        self.assertTrue(any("günde en fazla 2" in e for e in result.errors), result.errors)

    def test_session_mapping_uses_period_times(self):
        periods = list(self.school.periods.order_by("index"))
        mapping = exams.session_periods(self.period, periods, 0)
        self.assertEqual(mapping[0], [1])  # 09:30–10:30 → 2. saat (09:30–10:20)
        self.assertEqual(mapping[1], [2, 3])  # 11:00–12:00 → 10:30–11:20 ve 11:30–12:20


@override_settings(TIMETABLE_RUN_IN_BACKGROUND=False)
class ExamViewTests(TestCase):
    def setUp(self):
        self.school, self.period = build_exam_school(classes=2)
        self.user = User.objects.create_superuser("admin", "a@example.com", "x")
        self.client.force_login(self.user)

    def test_pages_render_and_flow(self):
        r = self.client.get(reverse("timetable:exams", args=[self.school.pk]))
        self.assertContains(r, "2026 Güz Vize")
        r = self.client.post(reverse("timetable:exams", args=[self.school.pk]), {
            "name": "2026 Güz Final", "kind": "final", "start_date": "2027-01-04", "end_date": "2027-01-15",
            "sessions_text": "13:30, 09:30", "max_per_day": 2, "proctors_per_room": 1, "capacity_percent": 50,
            "exam_minutes": 60, "time_limit": 10, "weekday_list": ["0", "1", "2", "3", "4", "5"],
            "excluded_text": "06.01.2027"})
        final = ExamPeriod.objects.get(name="2026 Güz Final")
        self.assertRedirects(r, reverse("timetable:exam_period", args=[self.school.pk, final.pk]))
        self.assertEqual(final.sessions, ["09:30", "13:30"])
        self.assertEqual(final.weekdays, [0, 1, 2, 3, 4, 5])
        self.assertEqual(final.excluded_dates, ["2027-01-06"])
        self.assertNotIn(datetime.date(2027, 1, 6), final.dates())
        self.assertIn(datetime.date(2027, 1, 9), final.dates())  # cumartesi

        url = reverse("timetable:exam_period", args=[self.school.pk, self.period.pk])
        r = self.client.post(url, {"action": "from_lessons", "mode": "replace"})
        self.assertEqual(self.period.exams.count(), 3)
        r = self.client.get(url)
        self.assertContains(r, "Matematik")
        self.assertContains(r, "Derslerden oluştur")
        r = self.client.post(reverse("timetable:exam_solve", args=[self.school.pk, self.period.pk]))
        self.period.refresh_from_db()
        self.assertEqual(self.period.status, ExamPeriod.OK, self.period.report)
        results = reverse("timetable:exam_results", args=[self.school.pk, self.period.pk])
        for view in ("takvim", "siniflar", "gozetmenler", "salonlar"):
            r = self.client.get(results + f"?gorunum={view}")
            self.assertEqual(r.status_code, 200, view)
            self.assertContains(r, "Matematik")
        for kind in ("takvim", "gozetmenler", "siniflar"):
            r = self.client.get(reverse("timetable:exam_print", args=[self.school.pk, self.period.pk]) + f"?tur={kind}")
            self.assertContains(r, "window.print()")
        r = self.client.get(reverse("timetable:exam_proctors", args=[self.school.pk, self.period.pk]))
        self.assertContains(r, "Öğretmen0")
        # gözetmen: muafiyet ve gelemeyeceği oturumlar
        t0 = Teacher.objects.get(school=self.school, first_name="Öğretmen0")
        proctor_url = reverse("timetable:exam_proctors", args=[self.school.pk, self.period.pk])
        self.client.post(proctor_url, {"action": "exempt", "exempt": [t0.pk]})
        self.assertEqual(list(self.period.proctor_exempt.all()), [t0])
        tuesday = (MONDAY + datetime.timedelta(days=1)).isoformat()
        self.client.post(proctor_url, {"action": "blocks", "teacher": t0.pk,
                                       "block": [f"{tuesday}#0", f"{tuesday}#1", f"{tuesday}#2", f"{MONDAY.isoformat()}#1"]})
        self.period.refresh_from_db()
        self.assertEqual(self.period.proctor_blocks[str(t0.pk)], [f"{MONDAY.isoformat()}#1", tuesday])
        self.assertIn((MONDAY, 1), exams.proctor_blocks(self.period)[t0.pk])
        self.assertContains(self.client.get(proctor_url + f"?ogretmen={t0.pk}"), "gelemeyeceği oturumlar")
        # ayarlar sekmesi
        r = self.client.get(url + "?sekme=ayarlar")
        self.assertContains(r, "Dönem ayarları")

        # elle taşıma
        exam = Exam.objects.get(period=self.period, subject__name="Fizik")
        edit = reverse("timetable:exam_edit", args=[self.school.pk, self.period.pk, exam.pk])
        self.assertEqual(self.client.get(edit).status_code, 200)
        room = Classroom.objects.filter(school=self.school).order_by("-capacity").first()
        target = MONDAY + datetime.timedelta(days=4)
        r = self.client.post(edit, {"action": "place", "date": target.isoformat(), "session": 2, "rooms": [room.pk]})
        pl = ExamPlacement.objects.get(exam=exam)
        self.assertEqual((pl.date, pl.session), (target, 2))
        self.assertEqual(list(pl.rooms.values_list("classroom_id", flat=True)), [room.pk])
        # gözetmen değiştirme: çakışan öğretmen seçilirse uyarı verilir ama kaydedilir
        er = pl.rooms.get()
        teacher = Teacher.objects.get(school=self.school, first_name="Öğretmen5")
        r = self.client.post(edit, {"action": "proctors", f"proctors-{er.pk}": [teacher.pk]}, follow=True)
        self.assertEqual(list(er.proctors.values_list("pk", flat=True)), [teacher.pk])

    def test_excel_export_opens(self):
        from openpyxl import load_workbook

        exams.generate_from_lessons(self.period)
        exams.start(self.period)
        r = self.client.get(reverse("timetable:exam_excel", args=[self.school.pk, self.period.pk]))
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Takvim", "Sınıflar", "Gözetmenler", "Salonlar"])
        values = [c.value for row in wb["Takvim"].iter_rows() for c in row]
        self.assertIn("Matematik", values)

    def test_viewer_cannot_edit(self):
        viewer = User.objects.create_user("viewer", password="x")
        Membership.objects.create(user=viewer, institution=self.school.institution, permissions=[Membership.VIEW])
        self.client.force_login(viewer)
        self.assertEqual(self.client.get(reverse("timetable:exams", args=[self.school.pk])).status_code, 200)
        r = self.client.post(reverse("timetable:exam_period", args=[self.school.pk, self.period.pk]),
                             {"action": "from_lessons", "mode": "replace"})
        self.assertEqual(r.status_code, 403)
        self.assertFalse(self.period.exams.exists())

    def test_other_school_period_is_404(self):
        other, other_period = build_exam_school()
        r = self.client.get(reverse("timetable:exam_period", args=[self.school.pk, other_period.pk]))
        self.assertEqual(r.status_code, 404)


class ExamAccessTests(TestCase):
    NAMES = ["exams", "exam_period", "exam_solve", "exam_results", "exam_edit", "exam_proctors", "exam_print", "exam_excel"]

    def test_routes_and_labels(self):
        for name in self.NAMES:
            self.assertIn(name, access.SCHOOL_ROUTES, name)
            self.assertIn(name, access.ACTION_LABELS, name)

    def test_ribbon_entry(self):
        from timetable.templatetags import navigation

        labels = [item[0] for _key, _title, items in navigation.NAV_GROUPS for item in items]
        self.assertIn("Sınav programı", labels)

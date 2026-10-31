"""ÜYBS aktarım şablonu, hafta arşivi, TC ve vekâlet işleme; e-posta bildirimleri."""

import datetime
from io import BytesIO

from django.contrib.auth.models import User
from django.core import mail
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from timetable import notify, refs, uybs
from timetable.models import (
    Absence, Card, Lesson, Period, School, SchoolClass, Subject, Substitution, Teacher, Timetable, WeekArchive,
)

MONDAY = datetime.date(2026, 11, 2)


class UybsTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name="MYO", days_per_week=5)
        for i in range(6):
            Period.objects.create(school=self.school, index=i, name=f"{i + 1}",
                                  start_time=datetime.time(8 + i, 20), end_time=datetime.time(9 + i, 0))
        self.cls = SchoolClass.objects.create(school=self.school, name="1.KISIM", short="1SVN")
        self.ali = Teacher.objects.create(school=self.school, first_name="Ali", last_name="Kaya", tc_no="22556185970", email="ali@x.tr")
        self.veli = Teacher.objects.create(school=self.school, first_name="Veli", last_name="Can", tc_no="59833013350", email="veli@x.tr")
        self.lesson = Lesson.objects.create(
            school=self.school, subject=Subject.objects.create(school=self.school, name="Askerî Tarih (U)", short="ASK203"),
            periods_per_week=2, card_length=2,
        )
        self.lesson.classes.set([self.cls])
        self.lesson.teachers.set([self.ali])
        self.timetable = Timetable.objects.create(school=self.school, name="Güz", status=Timetable.OK)
        self.card = Card.objects.create(timetable=self.timetable, lesson=self.lesson, length=2, day=0, period=3)
        School.objects.filter(pk=self.school.pk).update(active_timetable=self.timetable)
        self.school.refresh_from_db()
        refs.auto_number(self.school, self.timetable, "1001")
        self.user = User.objects.create_superuser("admin", "a@x.tr", "pw")
        self.client.force_login(self.user)

    def _sheet(self, response):
        return load_workbook(BytesIO(response.content))["UYBS_Ders_Aktarim"]

    def test_template_columns(self):
        self.school.uybs_settings = {"prefixes": {"ASK": "AB"}}
        self.school.save()
        url = reverse("timetable:uybs_export", args=[self.timetable.pk])
        ws = self._sheet(self.client.get(url))
        self.assertEqual(ws["B1"].value, "KISIM NO")
        self.assertEqual(ws["A1"].value, "ÜYBS REF")
        row = [ws.cell(row=2, column=c).value for c in range(1, 16)]
        self.assertEqual(row, [1001, 1, "ASK203", 0, "11:20", "12:00", "KSM", 1, 1, 1, "AB", "AB", 11, "22556185970", None])
        self.assertEqual(ws.cell(row=3, column=1).value, 1002)
        self.assertEqual(ws.cell(row=3, column=5).value, "12:20")

    def test_second_teacher_tc_and_ref_column_p(self):
        self.lesson.teachers.add(self.veli)
        self.school.uybs_settings = {"ref_col": "P"}
        self.school.save()
        ws = self._sheet(self.client.get(reverse("timetable:uybs_export", args=[self.timetable.pk])))
        self.assertIsNone(ws["A1"].value)
        self.assertEqual(ws["P1"].value, "ÜYBS REF")
        self.assertEqual({ws["N2"].value, ws["O2"].value}, {"22556185970", "59833013350"})
        self.assertEqual(ws["K2"].value, "ASK")  # eşleme yoksa ilk üç harf
        self.assertEqual(ws["P2"].value, 1001)

    def test_week_archive_keeps_past_week(self):
        url = reverse("timetable:uybs_export", args=[self.timetable.pk])
        self.client.get(url, {"hafta": "2026-11-04"})  # Çarşamba → Pazartesi'ye çekilir
        self.assertTrue(WeekArchive.objects.filter(school=self.school, week_start=MONDAY).exists())
        # Sonraki hafta ders 1–2. saate taşındı
        self.card.period = 0
        self.card.save()
        old = self._sheet(self.client.get(url, {"hafta": "2026-11-02"}))
        new = self._sheet(self.client.get(url, {"hafta": "2026-11-09"}))
        self.assertEqual((old["A2"].value, old["E2"].value), (1001, "11:20"))
        self.assertEqual((new["A2"].value, new["E2"].value), (1001, "08:20"))
        # Arşivi silince güncel programdan çıkar
        self.client.post(reverse("timetable:uybs_archive", args=[self.timetable.pk]), {"hafta": "2026-11-02", "action": "delete"})
        again = self._sheet(self.client.get(url, {"hafta": "2026-11-02"}))
        self.assertEqual(again["E2"].value, "08:20")

    def test_report_page_shows_archive_note(self):
        uybs.archive_week(self.timetable, MONDAY)
        response = self.client.get(reverse("timetable:report_render", args=[self.timetable.pk]), {"tur": "uybs", "hafta": "2026-11-02"})
        self.assertContains(response, "Arşivden")

    def test_reports_page_outputs_use_template(self):
        """Raporlar sayfasındaki ön izleme, Excel ve PDF de ÜYBS şablonu biçiminde (ayrı bir liste değil)."""
        self.school.uybs_settings = {"prefixes": {"ASK": "AB"}}
        self.school.save()
        params = {"tur": "uybs"}
        page = self.client.get(reverse("timetable:report_render", args=[self.timetable.pk]), params)
        self.assertContains(page, "KSM YAZACAK")
        self.assertContains(page, "<td>AB</td>")
        self.assertContains(page, "Referans kontrol listesi")
        ws = self._sheet(self.client.get(reverse("timetable:report_excel", args=[self.timetable.pk]), params))
        self.assertEqual((ws["A1"].value, ws["B1"].value), ("ÜYBS REF", "KISIM NO"))
        self.assertEqual([ws.cell(row=2, column=c).value for c in range(1, 16)],
                         [1001, 1, "ASK203", 0, "11:20", "12:00", "KSM", 1, 1, 1, "AB", "AB", 11, "22556185970", None])
        pdf = self.client.get(reverse("timetable:report_pdf", args=[self.timetable.pk]), params)
        self.assertEqual((pdf.status_code, pdf["Content-Type"]), (200, "application/pdf"))

    def test_substitution_uses_substitute_tc(self):
        absence = Absence.objects.create(school=self.school, date=MONDAY, teacher=self.ali)
        Substitution.objects.create(absence=absence, card=self.card, period=3, substitute=self.veli)
        rows, _ = uybs.week_rows(self.timetable, MONDAY)
        self.assertEqual(rows[0]["tcs"], ["59833013350"])
        self.assertIn("yerine vekil", rows[0]["sub"])
        self.assertEqual(rows[1]["tcs"], ["22556185970"])  # 2. saate vekil atanmadı
        rows, _ = uybs.week_rows(self.timetable, MONDAY, substitutions=False)
        self.assertEqual(rows[0]["tcs"], ["22556185970"])

    def test_class_number_default_and_override(self):
        self.assertEqual(refs.uybs_class_no(self.cls), "1")
        self.cls.uybs_no = "7"
        self.assertEqual(refs.uybs_class_no(self.cls), "7")

    def test_settings_page_and_tc_import(self):
        url = reverse("timetable:uybs_settings", args=[self.school.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        self.client.post(url, {"action": "template", "ksm": "KSM", "hep1": "1", "hep1b": "1", "hep11": "11", "ref_col": "A", "prefix-ASK": "AB"})
        self.school.refresh_from_db()
        self.assertEqual(self.school.uybs_settings["prefixes"], {"ASK": "AB"})
        self.client.post(url, {"action": "import_tc", "text": "Veli Can\t12345678901\tveli@okul.tr\nYok Biri\t1"})
        self.veli.refresh_from_db()
        self.assertEqual((self.veli.tc_no, self.veli.email), ("12345678901", "veli@okul.tr"))
        self.client.post(url, {"action": "people", f"class-{self.cls.pk}": "3", f"tc-{self.ali.pk}": "123", f"mail-{self.ali.pk}": "a@b.tr",
                               f"tc-{self.veli.pk}": "", f"mail-{self.veli.pk}": ""})
        self.cls.refresh_from_db()
        self.ali.refresh_from_db()
        self.assertEqual(self.cls.uybs_no, "3")
        self.assertEqual(self.ali.tc_no, "22556185970")  # geçersiz TC kaydedilmez


class NotifyTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name="Lise", days_per_week=5)
        for i in range(4):
            Period.objects.create(school=self.school, index=i, name=f"{i + 1}")
        cls = SchoolClass.objects.create(school=self.school, name="9A")
        self.ali = Teacher.objects.create(school=self.school, first_name="Ali", last_name="Kaya", email="ali@x.tr")
        self.ayse = Teacher.objects.create(school=self.school, first_name="Ayşe", last_name="Er", email="ayse@x.tr")
        self.lesson = Lesson.objects.create(school=self.school, subject=Subject.objects.create(school=self.school, name="Matematik", short="MAT"))
        self.lesson.classes.set([cls])
        self.lesson.teachers.set([self.ali])
        other = Lesson.objects.create(school=self.school, subject=Subject.objects.create(school=self.school, name="Fizik", short="FİZ"))
        other.classes.set([cls])
        other.teachers.set([self.ayse])
        self.old = Timetable.objects.create(school=self.school, name="v1", status=Timetable.OK)
        self.new = Timetable.objects.create(school=self.school, name="v2", status=Timetable.OK)
        Card.objects.create(timetable=self.old, lesson=self.lesson, day=0, period=0)
        Card.objects.create(timetable=self.new, lesson=self.lesson, day=1, period=2)
        Card.objects.create(timetable=self.old, lesson=other, day=2, period=0)
        Card.objects.create(timetable=self.new, lesson=other, day=2, period=0)
        School.objects.filter(pk=self.school.pk).update(active_timetable=self.old)
        self.client.force_login(User.objects.create_superuser("admin", "a@x.tr", "pw"))

    def test_changes_only_for_moved_teacher(self):
        diff = notify.changes(self.old, self.new)
        self.assertEqual(set(diff), {self.ali.id})
        self.assertEqual(len(diff[self.ali.id]), 2)

    def test_publish_sends_mail(self):
        self.client.post(reverse("timetable:publish", args=[self.new.pk]), {"notify": "1"})
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["ali@x.tr"])
        self.assertIn("Sal 3. saat", mail.outbox[0].body)
        self.client.post(reverse("timetable:publish", args=[self.old.pk]))  # kutu işaretsiz: e-posta yok
        self.assertEqual(len(mail.outbox), 1)

    def test_substitute_gets_mail(self):
        School.objects.filter(pk=self.school.pk).update(active_timetable=self.new)
        date = datetime.date(2026, 11, 3)  # Salı
        absence = Absence.objects.create(school=self.school, date=date, teacher=self.ali)
        card = self.new.cards.get(lesson=self.lesson)
        self.client.post(reverse("timetable:substitutions", args=[self.school.pk]), {
            "action": "save", "absence": absence.pk, "tarih": date.isoformat(), f"slot-{card.pk}-2": str(self.ayse.pk), "notify": "1",
        })
        self.assertEqual([m.to for m in mail.outbox], [["ayse@x.tr"]])
        self.assertIn("Ali Kaya yerine", mail.outbox[0].body)

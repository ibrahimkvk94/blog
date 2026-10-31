"""Program karşılaştırma karnesi: iki programın ölçütleri, kişi başı farklar, sayfa ve Excel."""

import io

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from timetable import access, program_compare
from timetable.input_diff import ASC_TIMETABLE
from timetable.models import Card, Lesson, Period, School, SchoolClass, Subject, Teacher, Timetable


class ProgramCompareTests(TestCase):
    def setUp(self):
        # 3 gün × 4 saat; saat bilgisi yok → geç saat = son saat (4.)
        self.school = School.objects.create(name="Okul", days_per_week=3)
        for i in range(4):
            Period.objects.create(school=self.school, index=i, name=str(i + 1))
        self.teacher = Teacher.objects.create(school=self.school, first_name="Ali", last_name="Kaya")
        self.cls = SchoolClass.objects.create(school=self.school, name="1.KISIM")
        # Ders başına günde bir kart kuralı ihlal edilmesin diye iki ders: MAT (2 saat), FİZ (1 saat)
        mat, fiz = [Lesson.objects.create(school=self.school, subject=Subject.objects.create(school=self.school, name=n, short=n),
                                          periods_per_week=hours, card_length=1) for n, hours in (("MAT", 2), ("FİZ", 1))]
        for lesson in (mat, fiz):
            lesson.teachers.add(self.teacher)
            lesson.classes.add(self.cls)
        # A: Pzt 1-2, Salı 1 → boşluk yok, 2 gün, 1 tek saatlik gün (+1 fazla gün)
        self.a = Timetable.objects.create(school=self.school, name="Üretilen", status=Timetable.OK)
        for lesson, day, period in ((mat, 0, 0), (fiz, 0, 1), (mat, 1, 0)):
            Card.objects.create(timetable=self.a, lesson=lesson, length=1, day=day, period=period)
        # B: Pzt 1 ve 4 (2 saat boşluk, 4. saat geç), bir kart bekliyor
        self.b = Timetable.objects.create(school=self.school, name=ASC_TIMETABLE, status=Timetable.OK)
        for lesson, day, period in ((mat, 0, 0), (fiz, 0, 3), (mat, None, None)):
            Card.objects.create(timetable=self.b, lesson=lesson, length=1, day=day, period=period)

    def test_summary_metrics_and_headline(self):
        data = program_compare.compare(self.a, self.b, threshold=1)
        m = {row["key"]: (row["a"], row["b"], row["better"]) for row in data["summary"]}
        self.assertEqual(m["hard"], (0, 0, None))
        self.assertEqual(m["unplaced"], (0, 1, "a"))
        self.assertEqual(m["score"], (7, 16, "a"))  # öğretmen 3+1 / 2·2+4, sınıf 3 (tek saatlik gün) / 2·2+4
        self.assertEqual(m["teacher_gaps"], (0, 2, "a"))
        self.assertEqual(m["gappy_teachers"], (0, 1, "a"))
        self.assertEqual(m["single_days"], (1, 0, "b"))
        self.assertEqual(m["avg_days"], (2, 1, "b"))
        self.assertEqual(m["max_days"], (2, 1, "b"))
        self.assertEqual(m["max_day"], (2, 2, None))
        self.assertEqual(m["late_hours"], (0, 1, "a"))
        self.assertEqual(m["class_gaps"], (0, 2, "a"))
        self.assertEqual(m["late_days"], (0, 0, None))
        wins = sum(1 for _, _, better in m.values() if better == "a")
        self.assertEqual(data["wins"], wins)
        self.assertEqual(data["verdict"], "a")
        self.assertEqual(data["headline"], f"Bu program {len(m)} ölçütün {program_compare.ordinal_locative(wins)} daha iyi")
        self.assertEqual(data["critical"]["key"], "unplaced")

    def test_per_teacher_and_class_rows(self):
        data = program_compare.compare(self.a, self.b)
        teacher = data["ogretmen"][0]
        self.assertEqual(teacher["id"], self.teacher.id)
        self.assertEqual(teacher["status"], "iyi")  # puan 4 < 8
        self.assertEqual((teacher["delta"]["score"], teacher["delta"]["gaps"], teacher["delta"]["days"], teacher["delta"]["late_hours"]), (-4, -2, 1, -1))
        cls = data["sinif"][0]
        self.assertEqual((cls["a"]["score"], cls["b"]["score"]), (3, 8))
        self.assertEqual(data["teacher_counts"], {"iyi": 1, "kotu": 0, "ayni": 0})
        # Ters yönde aynı satır kötüleşen olur
        reverse_data = program_compare.compare(self.b, self.a)
        self.assertEqual(reverse_data["ogretmen"][0]["status"], "kotu")
        self.assertEqual(reverse_data["verdict"], "b")

    def test_row_missing_on_one_side_counts_as_zero(self):
        other = Teacher.objects.create(school=self.school, first_name="Can", last_name="Ak")
        subject = Subject.objects.create(school=self.school, name="Fizik", short="FİZ")
        lesson = Lesson.objects.create(school=self.school, subject=subject, periods_per_week=1, card_length=1)
        lesson.teachers.add(other)
        lesson.classes.add(self.cls)
        Card.objects.create(timetable=self.a, lesson=lesson, length=1, day=2, period=1)
        data = program_compare.compare(self.a, self.b)
        row = next(p for p in data["ogretmen"] if p["id"] == other.id)
        self.assertEqual((row["a"]["hours"], row["b"]["hours"]), (1, 0))

    def test_ordinal_suffixes(self):
        self.assertEqual([program_compare.ordinal_locative(n) for n in (1, 2, 3, 5, 6, 7, 9, 10, 13, 20)],
                         ["1'inde", "2'sinde", "3'ünde", "5'inde", "6'sında", "7'sinde", "9'unda", "10'unda", "13'ünde", "20'sinde"])

    def test_page_defaults_to_asc_program_and_excel(self):
        self.client.force_login(User.objects.create_superuser("admin", "a@x.tr", "pw"))
        url = reverse("timetable:program_compare", args=[self.a.pk])
        page = self.client.get(url)
        self.assertEqual(page.context["other"], self.b)  # aSc'den gelen program önerilir
        self.assertContains(page, "ölçütün")
        self.assertContains(page, "Ali Kaya")
        self.assertContains(page, "Yalnız kötüleşenler")
        self.assertContains(page, 'data-status="iyi"')
        response = self.client.get(url + f"?diger={self.b.pk}&excel=1")
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["Özet", "Öğretmenler", "Sınıflar"])
        self.assertEqual(wb["Öğretmenler"]["A2"].value, "Ali Kaya")
        # Kalite karnesinden ve menüden ulaşılır
        self.assertContains(self.client.get(reverse("timetable:quality", args=[self.a.pk])), url)

    def test_other_school_program_is_not_compared(self):
        foreign_school = School.objects.create(name="Başka", days_per_week=3)
        foreign = Timetable.objects.create(school=foreign_school, name="x", status=Timetable.OK)
        self.client.force_login(User.objects.create_superuser("admin", "a@x.tr", "pw"))
        page = self.client.get(reverse("timetable:program_compare", args=[self.a.pk]) + f"?diger={foreign.pk}")
        self.assertIsNone(page.context["data"])
        self.assertNotContains(page, "data-cmp-headline")
        self.assertContains(page, "Karşılaştırmak için bir program seçin")

    def test_access_tables(self):
        self.assertEqual(access.TIMETABLE_ROUTES["program_compare"], (access.M.VIEW, access.M.VIEW))
        self.assertIn("program_compare", access.ACTION_LABELS)

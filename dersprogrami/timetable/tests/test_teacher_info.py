"""Öğretmen bilgileri toplu düzenleme ve Excel'de branş / TC sütunları."""

from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from timetable import importer
from timetable.models import School, Subject, Teacher


class TeacherInfoTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name="MYO")
        self.ahmet = Teacher.objects.create(school=self.school, first_name="Ahmet", last_name="Er", short="AER", title="Öğr.Gör.")
        self.ayse = Teacher.objects.create(school=self.school, first_name="Ayşe", last_name="Ak", short="AAK", tc_no="12345678901")
        self.client.force_login(User.objects.create_superuser("admin", "a@x.tr", "pw"))
        self.url = reverse("timetable:teacher_info", args=[self.school.pk])

    def post(self, **values):
        data = {}
        for t in (self.ahmet, self.ayse):
            for name in ("department", "title", "branch", "tc_no"):
                data[f"t-{t.pk}-{name}"] = getattr(t, name)
        data.update(values)
        return self.client.post(self.url, data, follow=True)

    def test_page_lists_teachers(self):
        page = self.client.get(self.url)
        self.assertContains(page, "Ahmet Er")
        self.assertContains(page, f'name="t-{self.ayse.pk}-tc_no" value="12345678901"', html=False)
        self.assertEqual(page.context["missing_tc"], 1)
        self.assertContains(self.client.get(reverse("timetable:data_list", args=[self.school.pk, "ogretmenler"])), self.url)

    def test_saves_changed_fields(self):
        response = self.post(**{f"t-{self.ahmet.pk}-department": "  ELN ", f"t-{self.ahmet.pk}-branch": "Elektronik",
                                f"t-{self.ahmet.pk}-tc_no": "109 876 543 21", f"t-{self.ayse.pk}-branch": "Matematik"})
        self.assertContains(response, "2 öğretmenin bilgisi güncellendi")
        self.ahmet.refresh_from_db()
        self.ayse.refresh_from_db()
        self.assertEqual((self.ahmet.department, self.ahmet.branch, self.ahmet.tc_no, self.ahmet.title),
                         ("ELN", "Elektronik", "10987654321", "Öğr.Gör."))
        self.assertEqual((self.ayse.tc_no, self.ayse.branch), ("12345678901", "Matematik"))

    def test_bad_tc_saves_nothing(self):
        response = self.post(**{f"t-{self.ahmet.pk}-department": "ELN", f"t-{self.ayse.pk}-tc_no": "123"})
        self.assertContains(response, "11 haneli olmalı")
        self.ahmet.refresh_from_db()
        self.assertEqual(self.ahmet.department, "")

    def test_requires_data_permission(self):
        self.client.logout()
        viewer = User.objects.create_user("izleyici", password="pw")
        self.client.force_login(viewer)
        self.assertNotEqual(self.client.get(self.url).status_code, 200)


class ExcelBranchTcTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name="MYO")
        Subject.objects.create(school=self.school, name="Matematik", short="MAT")
        self.teacher = Teacher.objects.create(school=self.school, first_name="Ahmet", last_name="Er", short="AER")

    def workbook(self, header, row):
        wb = Workbook()
        ws = wb.active
        ws.title = "Öğretmenler"
        ws.append(importer.SHEETS["Öğretmenler"] + header)
        ws.append(["Ahmet", "Er", "AER", None, None, None] + row)
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def test_branch_and_tc_columns(self):
        importer.import_workbook(self.school, self.workbook(["Branş", "TC kimlik no", "Okutabildiği dersler"],
                                                            ["Elektronik", 10987654321, "MAT"]))
        self.teacher.refresh_from_db()
        self.assertEqual((self.teacher.branch, self.teacher.tc_no), ("Elektronik", "10987654321"))
        self.assertEqual([s.short for s in self.teacher.subjects.all()], ["MAT"])

    def test_legacy_branches_column_sets_subjects(self):
        importer.import_workbook(self.school, self.workbook(["Branşlar"], ["MAT"]))
        self.teacher.refresh_from_db()
        self.assertEqual([s.short for s in self.teacher.subjects.all()], ["MAT"])
        self.assertEqual(self.teacher.branch, "")

    def test_export_has_branch_and_tc(self):
        self.teacher.branch, self.teacher.tc_no = "Elektronik", "10987654321"
        self.teacher.save()
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(importer.export_workbook(self.school)))
        ws = wb["Öğretmenler"]
        header = [c.value for c in ws[1]]
        values = dict(zip(header, [c.value for c in ws[2]]))
        self.assertEqual(values["Branş"], "Elektronik")
        self.assertEqual(str(values["TC kimlik no"]), "10987654321")


class SearchFoldTests(TestCase):
    def test_turkish_fold(self):
        from timetable.textsort import search_fold

        self.assertEqual(search_fold("İŞL103"), "isl103")
        self.assertEqual(search_fold("İBRAHİM Işık"), "ibrahim isik")
        self.assertEqual(search_fold("ibrahim"), "ibrahim")

    def test_search_attributes_use_fold(self):
        school = School.objects.create(name="MYO")
        Subject.objects.create(school=school, name="İktisat", short="İŞL103")
        from timetable.models import Lesson

        Lesson.objects.create(school=school, subject=school.subjects.get(), periods_per_week=2)
        self.client.force_login(User.objects.create_superuser("admin2", "b@x.tr", "pw"))
        page = self.client.get(reverse("timetable:teacher_assign", args=[school.pk]))
        self.assertContains(page, 'data-name="iktisat isl103"')

"""Programı Excel çalışma kitabına aktarır: sınıf, öğretmen ve derslik genel görünümleri ile kart listesi."""

from collections import defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .grid import KINDS, card_payload

SHEETS = [("sinif", "Sınıflar"), ("ogretmen", "Öğretmenler"), ("derslik", "Derslikler")]
THIN = Side(style="thin", color="CFD4DD")
THICK = Side(style="medium", color="5B6474")
HEADER_FILL = PatternFill("solid", fgColor="F4F5F8")
TITLE_FONT = Font(bold=True, size=13)
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _fill(color):
    return PatternFill("solid", fgColor=(color or "#FFFFFF").lstrip("#").upper())


def export(timetable):
    school = timetable.school
    periods = list(school.periods.all())
    days = school.day_names
    P = len(periods)
    cards = list(
        timetable.cards.select_related("lesson__subject", "classroom").prefetch_related(
            "lesson__classes", "lesson__teachers", "lesson__groups"
        )
    )
    wb = Workbook()
    wb.remove(wb.active)

    for kind, title in SHEETS:
        ws = wb.create_sheet(title)
        ws.cell(row=1, column=1, value=f"{school} – {timetable.name}").font = TITLE_FONT
        ws.cell(row=2, column=1, value=title).font = BOLD
        for d, day in enumerate(days):
            col = 2 + d * P
            ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + P - 1)
            cell = ws.cell(row=2, column=col, value=day)
            cell.font, cell.alignment, cell.fill = BOLD, CENTER, HEADER_FILL
            for j, period in enumerate(periods):
                c = ws.cell(row=3, column=col + j, value=f"{period.index + 1}\n{period.start_time:%H:%M}" if period.start_time else period.index + 1)
                c.alignment, c.fill = CENTER, HEADER_FILL
                c.border = Border(left=THICK if j == 0 else THIN, right=THIN, top=THIN, bottom=THIN)
        ws.cell(row=3, column=1).fill = HEADER_FILL

        payloads = [card_payload(c, kind) for c in cards if c.is_placed]
        by_row = defaultdict(lambda: defaultdict(list))
        for p in payloads:
            for row in p["rows"]:
                for i in range(p["length"]):
                    by_row[row][(p["day"], p["period"] + i)].append(p)

        for r, obj in enumerate(getattr(school, KINDS[kind]).all(), start=4):
            label = f"{obj.first_name} {obj.last_name}" if kind == "ogretmen" else str(obj)
            head = ws.cell(row=r, column=1, value=label)
            head.font, head.alignment = BOLD, LEFT
            head.border = Border(left=THIN, right=THICK, top=THIN, bottom=THIN)
            for d in range(len(days)):
                for j in range(P):
                    items = by_row[obj.id].get((d, j), [])
                    cell = ws.cell(row=r, column=2 + d * P + j)
                    if items:
                        cell.value = "\n".join(f"{i['subject']}\n{i['line2']}" if i["line2"] else i["subject"] for i in items)
                        cell.fill = _fill(items[0]["color"])
                    cell.alignment = CENTER
                    cell.font = Font(size=9)
                    cell.border = Border(left=THICK if j == 0 else THIN, right=THIN, top=THIN, bottom=THIN)
            ws.row_dimensions[r].height = 42
        ws.column_dimensions["A"].width = 24
        for col in range(2, 2 + len(days) * P):
            ws.column_dimensions[get_column_letter(col)].width = 10
        ws.row_dimensions[3].height = 30
        ws.freeze_panes = "B4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws = wb.create_sheet("Kartlar")
    headers = ["Ders", "Sınıf", "Grup", "Öğretmen", "Derslik", "Gün", "Saat", "Süre", "Kilitli"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = BOLD, HEADER_FILL
    for r, card in enumerate(cards, start=2):
        p = card_payload(card, "sinif")
        values = [
            p["subject_name"],
            p["classes"],
            p["groups"],
            p["teachers"],
            p["room"],
            days[card.day] if card.is_placed else "Yerleşmedi",
            card.period + 1 if card.is_placed else "",
            card.length,
            "Evet" if card.locked else "",
        ]
        for i, v in enumerate(values, start=1):
            ws.cell(row=r, column=i, value=v)
    for i, width in enumerate([24, 16, 16, 22, 14, 12, 7, 7, 8], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(cards) + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

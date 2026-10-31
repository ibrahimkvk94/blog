"""Sınav programının Excel çıktısı: Takvim, Sınıflar, Gözetmenler ve Salonlar sayfaları."""

from io import BytesIO

from openpyxl import Workbook

from . import exams
from .models import DAY_NAMES
from .report_excel import _header, _table


def _time(period, pl):
    times = period.session_times
    start = times[pl.session] if pl.session < len(times) else f"{pl.session + 1}. oturum"
    if pl.exam.duration > 1:
        start += f" ({pl.exam.duration} oturum)"
    return start


def _classes(pl):
    return ", ".join(c.short or c.name for c in pl.exam.classes.all())


def _rooms(pl, with_proctors=False):
    parts = []
    for r in pl.rooms.all():
        text = f"{r.classroom} ({r.seats})"
        if with_proctors and r.proctors.all():
            text += ": " + ", ".join(t.full_name for t in r.proctors.all())
        parts.append(text)
    return "\n".join(parts)


def build(period):
    rows = exams.placements(period)
    wb = Workbook()
    wb.remove(wb.active)
    info = f"{period.school} · {period.start_date:%d.%m.%Y} – {period.end_date:%d.%m.%Y}"

    ws = wb.create_sheet("Takvim")
    start = _header(ws, f"{period.name} – Sınav takvimi", info, span=8)
    _table(ws, start, ["Tarih", "Gün", "Saat", "Sınav", "Sınıflar", "Öğrenci", "Salonlar ve gözetmenler", "Dersin öğretmeni"], [
        [pl.date.strftime("%d.%m.%Y"), DAY_NAMES[pl.date.weekday()], _time(period, pl), pl.exam.title, _classes(pl),
         pl.exam.student_count, _rooms(pl, with_proctors=True), ", ".join(t.full_name for t in pl.exam.teachers.all())]
        for pl in rows
    ], widths=[12, 11, 12, 34, 24, 9, 50, 30])

    ws = wb.create_sheet("Sınıflar")
    start = _header(ws, f"{period.name} – Sınıfların sınavları", info, span=6)
    data = []
    for c, items in exams.by_class(period, rows):
        for pl, rooms in items:
            data.append([c.short or c.name, pl.date.strftime("%d.%m.%Y"), DAY_NAMES[pl.date.weekday()], _time(period, pl),
                         pl.exam.title, ", ".join(str(r.classroom) for r in rooms)])
    _table(ws, start, ["Sınıf", "Tarih", "Gün", "Saat", "Sınav", "Salonlar"], data, widths=[14, 12, 11, 12, 34, 34])

    ws = wb.create_sheet("Gözetmenler")
    start = _header(ws, f"{period.name} – Gözetmen listesi", info, span=7)
    data = []
    for t, items in exams.by_teacher(period, rows):
        for k, (pl, room) in enumerate(items):
            data.append([t.full_name if k == 0 else "", len(items) if k == 0 else None, pl.date.strftime("%d.%m.%Y"),
                         DAY_NAMES[pl.date.weekday()], _time(period, pl), pl.exam.title, str(room.classroom)])
    _table(ws, start, ["Öğretmen", "Görev", "Tarih", "Gün", "Saat", "Sınav", "Salon"], data, widths=[28, 8, 12, 11, 12, 34, 16])

    ws = wb.create_sheet("Salonlar")
    start = _header(ws, f"{period.name} – Salonlar", info, span=7)
    data = []
    for room, items in exams.by_room(period, rows):
        for pl, r in items:
            data.append([str(room), pl.date.strftime("%d.%m.%Y"), DAY_NAMES[pl.date.weekday()], _time(period, pl),
                         pl.exam.title, r.seats, ", ".join(t.full_name for t in r.proctors.all())])
    _table(ws, start, ["Salon", "Tarih", "Gün", "Saat", "Sınav", "Öğrenci", "Gözetmenler"], data, widths=[16, 12, 11, 12, 34, 9, 40])

    out = BytesIO()
    wb.save(out)
    return out.getvalue()

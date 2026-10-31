"""Excel ile toplu veri aktarımı.

Okulun verisi beş sayfalı bir çalışma kitabına aktarılır; aynı dosya şablon olarak da kullanılır.
Geri yüklemede kayıtlar adlarına göre eşleştirilir: varsa güncellenir, yoksa eklenir. Herhangi bir
satırda hata olursa hiçbir değişiklik kaydedilmez ve hatalar satır numarasıyla bildirilir.
"""

import re
from collections import defaultdict
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Building, Classroom, Division, Group, Lesson, SchoolClass, Subject, Teacher

PALETTE = ["#dfe7fd", "#fde2e4", "#e2ece9", "#fff1e6", "#e9f5db", "#f0efeb", "#cddafd", "#fad2e1", "#bee1e6", "#eae4e9"]

SHEETS = {
    "Ders adları": ["Ad", "Kısa ad", "Renk"],
    "Derslikler": ["Ad", "Kısa ad", "Kapasite", "Bina"],
    "Sınıflar": ["Ad", "Seviye", "Sınıf dersliği", "Bölünmeler"],
    "Öğretmenler": ["Ad", "Soyad", "Kısa ad", "Günde en fazla ders", "Haftada en fazla gün", "Haftada en fazla boşluk"],
    "Dersler": ["Ders adı", "Sınıflar", "Gruplar", "Öğretmenler", "Derslikler", "Haftalık saat", "Kart uzunluğu", "Günde en fazla"],
}

# Sonradan eklenen sütunlar: eski dosyalarda yoksa ilgili alan değiştirilmez.
OPTIONAL = {
    "Ders adları": ["Varsayılan derslikler"],
    "Sınıflar": ["Öğrenci sayısı", "Sınıf öğretmeni"],
    "Öğretmenler": ["Haftalık ders yükü", "Okutabildiği dersler", "Öğretmen dersliği", "Bölüm", "Ünvan / Rütbe", "Branş", "TC kimlik no"],
    "Dersler": ["Kart dağılımı", "Derslik sayısı"],
}
# Şablonda yer almayan ama okunmaya devam eden eski sütunlar (ünvan ve rütbe önceden ayrıydı).
LEGACY = {"Öğretmenler": ["Ünvan", "Rütbe", "Branşlar"]}  # "Branşlar": okutabildiği derslerin eski adı

NOTES = {
    "Sınıflar": "Bölünmeler: “Yabancı dil 2: Almanca, Fransızca; Beden: Kız, Erkek” biçiminde.",
    "Öğretmenler": "Okutabildiği dersler: ders adları ya da kısa adları virgülle. Branş: düz metin (ör. Bilgisayar). "
    "Haftalık ders yükü: ör. 21 (maaş karşılığı + ek ders).",
    "Dersler": "Sınıflar, gruplar, öğretmenler ve derslikler virgülle ayrılır. Öğretmen kısa adı ya da ad soyadı yazılabilir. "
    "Grup birden çok sınıfta aynı adı taşıyorsa yalnızca grup adı yeterlidir (ör. 9-A, 9-B + Almanca). "
    "Sınıfsız (yalnızca öğretmenli) satır toplantı/hazırlık saati olur. Kart dağılımı: ör. 2+2+1.",
}

HEADER_FILL = PatternFill("solid", fgColor="EEF0FF")


class ImportFailed(Exception):
    def __init__(self, errors):
        super().__init__("\n".join(errors))
        self.errors = errors


def _key(value):
    return str(value or "").strip().casefold()


def _split(value, sep=","):
    return [part.strip() for part in str(value or "").split(sep) if part.strip()]


def _int(value):
    if value in (None, ""):
        return None
    return int(float(str(value).replace(",", ".")))


# --- dışa aktarma ------------------------------------------------------------------------------


def export_workbook(school):
    wb = Workbook()
    wb.remove(wb.active)
    rows = {
        "Ders adları": [
            [s.name, s.short, s.color, ", ".join(str(r) for r in s.classrooms.all())]
            for s in school.subjects.prefetch_related("classrooms").order_by("name")
        ],
        "Derslikler": [[r.name, r.short, r.capacity, r.building.name if r.building else ""] for r in school.classrooms.select_related("building").order_by("name")],
        "Sınıflar": [
            [
                c.name,
                c.grade,
                str(c.home_classroom) if c.home_classroom else "",
                "; ".join(f"{d.name}: {', '.join(g.name for g in d.groups.all())}" for d in c.divisions.all()),
                c.student_count,
                (c.class_teacher.short or c.class_teacher.full_name) if c.class_teacher else "",
            ]
            for c in school.classes.select_related("home_classroom", "class_teacher").prefetch_related("divisions__groups")
        ],
        "Öğretmenler": [
            [
                t.first_name, t.last_name, t.short, t.max_periods_per_day, t.max_days_per_week, t.max_gaps_per_week,
                t.target_periods, ", ".join(sub.short or sub.name for sub in t.subjects.all()),
                str(t.home_classroom) if t.home_classroom else "",
                t.department, t.title, t.branch, t.tc_no,
            ]
            for t in school.teachers.select_related("home_classroom").prefetch_related("subjects")
        ],
        "Dersler": [
            [
                lesson.subject.name,
                ", ".join(c.name for c in lesson.classes.all()),
                ", ".join(sorted({g.name for g in lesson.groups.all()})),
                ", ".join(t.short or t.full_name for t in lesson.teachers.all()),
                ", ".join(str(r) for r in lesson.classrooms.all()),
                lesson.periods_per_week,
                lesson.card_length,
                lesson.max_per_day,
                lesson.split,
                lesson.rooms_needed if lesson.rooms_needed > 1 else None,
            ]
            for lesson in school.lessons.select_related("subject").prefetch_related("classes", "groups", "teachers", "classrooms").order_by("subject__name", "id")
        ],
    }
    for title, headers in SHEETS.items():
        headers = headers + OPTIONAL.get(title, [])
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), HEADER_FILL
        for row in rows[title]:
            ws.append(row)
        for i, header in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(14, len(header) + 4)
        ws.freeze_panes = "A2"
        if title == "Dersler":
            ws.column_dimensions["B"].width = ws.column_dimensions["D"].width = 22
            validation = DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=True)
            ws.add_data_validation(validation)
            validation.add("G2:G2000")
        if title == "Sınıflar":
            ws.column_dimensions["D"].width = 48
    info = wb.create_sheet("Açıklama", 0)
    info.column_dimensions["A"].width = 110
    lines = [
        f"{school} – veri aktarım dosyası",
        "",
        "Bu dosyayı düzenleyip uygulamada Veri → Excel aktarımı sayfasından yükleyin.",
        "Kayıtlar adlarına göre eşleştirilir: var olan kayıt güncellenir, olmayan eklenir. Sayfa adlarını ve başlıkları değiştirmeyin.",
        "Bir satırda hata varsa hiçbir değişiklik kaydedilmez; hatalar satır numarasıyla gösterilir.",
        "",
        *(f"{sheet}: {note}" for sheet, note in NOTES.items()),
    ]
    for i, line in enumerate(lines, start=1):
        info.cell(row=i, column=1, value=line).font = Font(bold=i == 1, size=13 if i == 1 else 11)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- içe aktarma -------------------------------------------------------------------------------


class _Importer:
    def __init__(self, school, delete_missing_lessons=False):
        self.school = school
        self.delete_missing = delete_missing_lessons
        self.errors = []
        self.counts = defaultdict(lambda: {"eklenen": 0, "güncellenen": 0})

    def error(self, sheet, row, message):
        self.errors.append(f"{sheet}, satır {row}: {message}")

    def rows(self, wb, title):
        if title not in wb.sheetnames:
            return []
        ws = wb[title]
        headers = [_key(c.value) for c in ws[1]]
        expected = [_key(h) for h in SHEETS[title]]
        missing = [h for h, k in zip(SHEETS[title], expected) if k not in headers]
        if missing:
            self.errors.append(f"{title}: eksik sütun başlığı: {', '.join(missing)}")
            return []
        names = SHEETS[title] + [h for h in OPTIONAL.get(title, []) + LEGACY.get(title, []) if _key(h) in headers]
        index = {_key(h): headers.index(_key(h)) for h in names}
        result = []
        for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            values = {h: (row[index[_key(h)]] if index[_key(h)] < len(row) else None) for h in names}
            if all(v in (None, "") for v in values.values()):
                continue
            result.append((number, values))
        return result

    def count(self, title, created):
        self.counts[title]["eklenen" if created else "güncellenen"] += 1

    # Arama tabloları; her sayfadan sonra yenilenir.
    def refresh(self):
        school = self.school
        self.subjects = {}
        for s in school.subjects.all():
            self.subjects[_key(s.name)] = s
            if s.short:
                self.subjects.setdefault(_key(s.short), s)
        self.rooms = {}
        for r in school.classrooms.all():
            self.rooms[_key(r.name)] = r
            if r.short:
                self.rooms.setdefault(_key(r.short), r)
        self.classes = {_key(c.name): c for c in school.classes.all()}
        self.teachers = {}
        for t in school.teachers.all():
            self.teachers[_key(t.full_name)] = t
            if t.short:
                self.teachers.setdefault(_key(t.short), t)

    def lookup(self, title, n, table, value, noun):
        """Virgüllü adları kayıtlara çevirir; bulunamayan varsa hata yazar ve None döner."""
        found, missing = [], []
        for name in _split(value):
            obj = table.get(_key(name))
            (found if obj is not None else missing).append(obj if obj is not None else name)
        if missing:
            self.error(title, n, ", ".join(f"“{m}”" for m in missing) + f" {noun} bulunamadı.")
            return None
        return found

    def subjects_sheet(self, wb):
        title = "Ders adları"
        self.refresh()
        for i, (n, v) in enumerate(self.rows(wb, title)):
            name = str(v["Ad"] or "").strip()
            if not name:
                self.error(title, n, "Ad boş olamaz.")
                continue
            color = str(v["Renk"] or "").strip()
            subject, created = Subject.objects.get_or_create(
                school=self.school, name=name, defaults={"color": PALETTE[i % len(PALETTE)]}
            )
            subject.short = str(v["Kısa ad"] or "").strip()[:20]
            if color.startswith("#") and len(color) == 7:
                subject.color = color
            subject.save()
            if "Varsayılan derslikler" in v:
                rooms = self.lookup(title, n, self.rooms, v["Varsayılan derslikler"], "dersliği")
                if rooms is None:
                    continue
                subject.classrooms.set(rooms)
            self.count(title, created)

    def rooms_sheet(self, wb):
        title = "Derslikler"
        for n, v in self.rows(wb, title):
            name = str(v["Ad"] or "").strip()
            if not name:
                self.error(title, n, "Ad boş olamaz.")
                continue
            try:
                capacity = _int(v["Kapasite"])
            except ValueError:
                self.error(title, n, f"Kapasite sayı olmalı: “{v['Kapasite']}”.")
                continue
            building = None
            if v["Bina"]:
                building, _ = Building.objects.get_or_create(school=self.school, name=str(v["Bina"]).strip())
            room, created = Classroom.objects.get_or_create(school=self.school, name=name)
            room.short = str(v["Kısa ad"] or "").strip()[:20]
            room.capacity, room.building = capacity, building
            room.save()
            self.count(title, created)

    def classes_sheet(self, wb):
        title = "Sınıflar"
        self.refresh()
        for n, v in self.rows(wb, title):
            name = str(v["Ad"] or "").strip()
            if not name:
                self.error(title, n, "Ad boş olamaz.")
                continue
            home = None
            if v["Sınıf dersliği"]:
                home = self.rooms.get(_key(v["Sınıf dersliği"]))
                if home is None:
                    self.error(title, n, f"“{v['Sınıf dersliği']}” dersliği bulunamadı (önce Derslikler sayfasına ekleyin).")
                    continue
            try:
                grade = _int(v["Seviye"])
            except ValueError:
                self.error(title, n, f"Seviye sayı olmalı: “{v['Seviye']}”.")
                continue
            if "Öğrenci sayısı" in v:
                try:
                    students = _int(v["Öğrenci sayısı"])
                except ValueError:
                    self.error(title, n, f"Öğrenci sayısı sayı olmalı: “{v['Öğrenci sayısı']}”.")
                    continue
            class_teacher = None
            if v.get("Sınıf öğretmeni"):
                class_teacher = self.teachers.get(_key(v["Sınıf öğretmeni"]))
                if class_teacher is None:
                    self.error(title, n, f"“{v['Sınıf öğretmeni']}” öğretmeni bulunamadı.")
                    continue
            cls, created = SchoolClass.objects.get_or_create(school=self.school, name=name)
            cls.grade, cls.home_classroom = grade, home
            if "Öğrenci sayısı" in v:
                cls.student_count = students
            if "Sınıf öğretmeni" in v:
                cls.class_teacher = class_teacher
            cls.save()
            for part in _split(v["Bölünmeler"], ";"):
                division_name, _, groups = part.partition(":")
                group_names = _split(groups)
                if not division_name.strip() or len(group_names) < 2:
                    self.error(title, n, f"Bölünme “{part}” en az iki grup içermeli (ör. “Dil: Almanca, Fransızca”).")
                    continue
                division, _ = Division.objects.get_or_create(school_class=cls, name=division_name.strip())
                for group_name in group_names:
                    Group.objects.get_or_create(division=division, name=group_name)
            self.count(title, created)

    def teachers_sheet(self, wb):
        title = "Öğretmenler"
        self.refresh()
        for n, v in self.rows(wb, title):
            first, last = str(v["Ad"] or "").strip(), str(v["Soyad"] or "").strip()
            if not first:
                self.error(title, n, "Ad boş olamaz.")
                continue
            try:
                limits = [_int(v[h]) for h in ("Günde en fazla ders", "Haftada en fazla gün", "Haftada en fazla boşluk")]
            except ValueError:
                self.error(title, n, "Sınır sütunları sayı olmalı.")
                continue
            try:
                target = _int(v.get("Haftalık ders yükü"))
            except ValueError:
                self.error(title, n, "Haftalık ders yükü sayı olmalı.")
                continue
            subject_column = "Okutabildiği dersler" if "Okutabildiği dersler" in v else "Branşlar"
            subjects = self.lookup(title, n, self.subjects, v.get(subject_column), "ders adı")
            home = None
            if v.get("Öğretmen dersliği"):
                home = self.rooms.get(_key(v["Öğretmen dersliği"]))
                if home is None:
                    self.error(title, n, f"“{v['Öğretmen dersliği']}” dersliği bulunamadı.")
                    continue
            if subjects is None:
                continue
            teacher = Teacher.objects.filter(school=self.school, first_name=first, last_name=last).first()
            created = teacher is None
            if created:
                teacher = Teacher(school=self.school, first_name=first, last_name=last)
            teacher.short = str(v["Kısa ad"] or "").strip()[:20]
            teacher.max_periods_per_day, teacher.max_days_per_week, teacher.max_gaps_per_week = limits
            if "Haftalık ders yükü" in v:
                teacher.target_periods = target
            if "Öğretmen dersliği" in v:
                teacher.home_classroom = home
            if "Bölüm" in v:
                teacher.department = str(v["Bölüm"] or "").strip()[:60]
            # Ünvan ve rütbe tek alandır: "Ünvan / Rütbe" sütunu ya da ayrı "Ünvan", "Rütbe" sütunları birleştirilir.
            title_columns = [c for c in ("Ünvan / Rütbe", "Ünvan", "Rütbe") if c in v]
            if title_columns:
                teacher.title = " ".join(filter(None, (str(v[c] or "").strip() for c in title_columns)))[:100]
            if "Branş" in v:
                teacher.branch = str(v["Branş"] or "").strip()[:100]
            if "TC kimlik no" in v:
                tc = re.sub(r"\D", "", str(v["TC kimlik no"] or "").split(".")[0])
                if tc and len(tc) != 11:
                    self.error(title, n, f"TC kimlik no 11 haneli olmalı ({v['TC kimlik no']}).")
                    continue
                teacher.tc_no = tc
            teacher.save()
            if subject_column in v:
                teacher.subjects.set(subjects)
            self.count(title, created)

    def lessons_sheet(self, wb):
        title = "Dersler"
        self.refresh()
        existing = {}
        for lesson in self.school.lessons.prefetch_related("classes", "groups"):
            key = (lesson.subject_id, frozenset(c.id for c in lesson.classes.all()), frozenset(g.id for g in lesson.groups.all()))
            existing.setdefault(key, lesson)
        seen = set()
        rows = self.rows(wb, title)
        for n, v in rows:
            problems = []
            subject = self.subjects.get(_key(v["Ders adı"]))
            if subject is None:
                problems.append(f"“{v['Ders adı']}” ders adı bulunamadı")
            classes = []
            for name in _split(v["Sınıflar"]):
                cls = self.classes.get(_key(name))
                if cls is None:
                    problems.append(f"“{name}” sınıfı bulunamadı")
                else:
                    classes.append(cls)
            if not _split(v["Sınıflar"]) and not _split(v["Öğretmenler"]):
                problems.append("en az bir sınıf ya da öğretmen yazılmalı")
            groups = []
            for name in _split(v["Gruplar"]):
                found = [
                    g for g in Group.objects.filter(division__school_class__in=classes)
                    if _key(g.name) == _key(name) or _key(f"{g.division.school_class.name} {g.name}") == _key(name)
                ]
                if not found:
                    problems.append(f"“{name}” grubu seçilen sınıflarda yok")
                groups += found
            teachers = []
            for name in _split(v["Öğretmenler"]):
                teacher = self.teachers.get(_key(name))
                if teacher is None:
                    problems.append(f"“{name}” öğretmeni bulunamadı")
                else:
                    teachers.append(teacher)
            rooms = []
            for name in _split(v["Derslikler"]):
                room = self.rooms.get(_key(name))
                if room is None:
                    problems.append(f"“{name}” dersliği bulunamadı")
                else:
                    rooms.append(room)
            try:
                hours = _int(v["Haftalık saat"])
                length = _int(v["Kart uzunluğu"]) or 1
                per_day = _int(v["Günde en fazla"])
            except ValueError:
                problems.append("saat sütunları sayı olmalı")
                hours = length = per_day = None
            if hours is not None and hours < 1:
                problems.append("haftalık saat en az 1 olmalı")
            if hours is None and not problems:
                problems.append("haftalık saat boş olamaz")
            if length is not None and not 1 <= length <= 4:
                problems.append("kart uzunluğu 1–4 arasında olmalı")
            split = str(v.get("Kart dağılımı") or "").strip()
            if split:
                parts = Lesson.parse_split(split)
                if not parts or (hours and sum(parts) != hours):
                    problems.append(f"kart dağılımı “{split}” geçersiz ya da toplamı haftalık saate eşit değil")
                else:
                    split = "+".join(str(p) for p in parts)
            try:
                rooms_needed = _int(v.get("Derslik sayısı")) or 1
            except ValueError:
                problems.append("derslik sayısı sayı olmalı")
                rooms_needed = 1
            if problems:
                self.error(title, n, "; ".join(problems) + ".")
                continue
            key = (subject.id, frozenset(c.id for c in classes), frozenset(g.id for g in groups))
            lesson = existing.get(key)
            created = lesson is None
            if created:
                lesson = Lesson.objects.create(school=self.school, subject=subject, periods_per_week=hours)
                lesson.classes.set(classes)
                lesson.groups.set(groups)
                existing[key] = lesson
            lesson.periods_per_week, lesson.card_length, lesson.max_per_day = hours, length, per_day
            if "Kart dağılımı" in v:
                lesson.split = split
            if "Derslik sayısı" in v:
                lesson.rooms_needed = rooms_needed
            lesson.save()
            lesson.teachers.set(teachers)
            lesson.classrooms.set(rooms)
            seen.add(lesson.pk)
            self.count(title, created)
        if self.delete_missing and rows:
            removed = self.school.lessons.exclude(pk__in=seen)
            self.counts[title]["silinen"] = removed.count()
            removed.delete()

    def run(self, data):
        try:
            wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
        except Exception:
            raise ImportFailed(["Dosya okunamadı. Uygulamadan indirilen .xlsx şablonunu kullanın."])
        if not any(title in wb.sheetnames for title in SHEETS):
            raise ImportFailed([f"Dosyada beklenen sayfalardan hiçbiri yok ({', '.join(SHEETS)})."])
        with transaction.atomic():
            self.rooms_sheet(wb)
            self.subjects_sheet(wb)
            self.teachers_sheet(wb)
            self.classes_sheet(wb)
            self.lessons_sheet(wb)
            if self.errors:
                transaction.set_rollback(True)
        if self.errors:
            raise ImportFailed(self.errors)
        return {title: dict(counts) for title, counts in self.counts.items()}


def import_workbook(school, data, delete_missing_lessons=False):
    """Dönüş: {sayfa: {"eklenen": n, "güncellenen": n[, "silinen": n]}}. Hata varsa ImportFailed."""
    return _Importer(school, delete_missing_lessons).run(data)

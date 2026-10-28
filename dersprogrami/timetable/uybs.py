"""ÜYBS ders aktarım şablonu ve haftalık arşiv.

Şablon (UYBS_Ders_Aktarim sayfası) her kısımdaki her ders saati için bir satırdır:
B kısım no, C ders kodu, D gün (Pazartesi 0 … Cuma 4), E başlangıç, F bitiş, G "KSM", H ders adında
"(U)" varsa 1, I/J "1", K/L ders kodu önekinin ÜYBS karşılığı, M "11", N/O öğretim elemanı TC.
Referans varsayılan olarak A sütununa yazılır.

Bir hafta aktarıldığında satırlar o haftanın arşivine yazılır; program sonradan değişse de o
haftanın raporu aynı kalır. Arşiv silinirse rapor güncel programdan yeniden üretilir.
"""

import datetime
import re
from io import BytesIO

from django.contrib import messages
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import refs
from .models import School, Teacher, Timetable, WeekArchive

DEFAULTS = {"ksm": "KSM", "hep1": "1", "hep1b": "1", "hep11": "11", "ref_col": "A", "prefixes": {}}
REF_COLUMNS = {"A": "A sütunu (ilk sütun)", "P": "P sütunu (son sütun)", "": "Yazılmasın"}
HEADERS = [
    "KISIM NO", "DERS KODU", "GÜN\nPazartesi 0\n...\nCuma 4", "BAŞLANGIÇ", "BİTİŞ", "KSM YAZACAK",
    "U / ORTAK\n(DERS ADI İÇİNDE (U) İFADESİ VARSA 1 YOKSA 0)", "HEP 1", "HEP 1",
    "DERS KODUNUN SOLDAN İLK 3 HARFİ", "DERS KODUNUN SOLDAN İLK 3 HARFİ", "HEP 11", "ÖĞRETİM ELEMANI 1 TC.",
    "ÖĞRETİM ELEMANI 2 TC\n(Birden fazla öğretmen giriyorsa 2. öğretim elemanı Tc. buraya)",
]
TC_RE = re.compile(r"\d{11}")


def settings_of(school):
    data = dict(DEFAULTS)
    data.update(school.uybs_settings or {})
    data["prefixes"] = dict(data.get("prefixes") or {})
    return data


def prefix_of(code):
    return (code or "").strip()[:3].upper()


def unit_code(code, settings):
    prefix = prefix_of(code)
    return settings["prefixes"].get(prefix) or prefix


def _number(value):
    """Sayı gibi duran sabitler Excel'e sayı olarak yazılır."""
    return int(value) if str(value).isdigit() else value


def monday(date):
    return date - datetime.timedelta(days=date.weekday())


# --- haftalık satırlar ve arşiv ---------------------------------------------------------------

def _dump(rows):
    return [{**r, "date": r["date"].isoformat() if r["date"] else None, "sort": list(r["sort"])} for r in rows]


def _load(rows):
    return [{**r, "date": datetime.date.fromisoformat(r["date"]) if r.get("date") else None} for r in rows]


def week_rows(timetable, week_start=None, class_ids=None, substitutions=True, save=False, user=None):
    """Dönüş: (satırlar, arşiv ya da None). Haftanın arşivi varsa satırlar arşivden gelir;
    save=True ise arşivsiz hafta güncel programdan üretilip arşivlenir."""
    school = timetable.school
    archive = WeekArchive.objects.filter(school=school, week_start=week_start).first() if week_start else None
    if archive is not None:
        rows = _load(archive.rows)
    else:
        rows = refs.report_rows(timetable, week_start, None, substitutions=substitutions)
        if save and week_start:
            archive = WeekArchive.objects.create(
                school=school, week_start=week_start, source=timetable.name, rows=_dump(rows), user=user
            )
    if class_ids:
        rows = [r for r in rows if r.get("class_id") in set(class_ids)]
    return rows, archive


def archive_week(timetable, week_start, user=None, substitutions=True):
    """Haftayı güncel programdan (yeniden) arşivler."""
    rows = refs.report_rows(timetable, week_start, None, substitutions=substitutions)
    archive, _ = WeekArchive.objects.update_or_create(
        school=timetable.school, week_start=week_start,
        defaults={"source": timetable.name, "rows": _dump(rows), "user": user},
    )
    return archive


# --- şablon ------------------------------------------------------------------------------------

def template_rows(school, rows):
    """Şablon satırları (yalnızca yerleşmiş saatler). Arşivde TC boşsa öğretmenin güncel TC'si alınır."""
    settings = settings_of(school)
    tc_now = dict(Teacher.objects.filter(school=school).values_list("id", "tc_no"))
    out = []
    for r in rows:
        if r.get("day_index") is None:
            continue
        tcs = [tc or tc_now.get(tid, "") for tc, tid in zip(r.get("tcs") or [], r.get("teacher_ids") or [])]
        tcs = [tc for tc in tcs if tc] + ["", ""]
        unit = unit_code(r["code"], settings)
        out.append({
            "ref": r["ref"],
            "cells": [
                _number(r.get("kisim_no") or r["class"]), r["code"], r["day_index"], r.get("start", ""), r.get("end", ""),
                settings["ksm"], 1 if "(U)" in (r["subject"] or "").upper() else 0,
                _number(settings["hep1"]), _number(settings["hep1b"]), unit, unit, _number(settings["hep11"]),
                tcs[0], tcs[1],
            ],
        })
    return out


def template_xlsx(school, rows):
    settings = settings_of(school)
    wb = Workbook()
    ws = wb.active
    ws.title = "UYBS_Ders_Aktarim"
    ref_col = settings["ref_col"]
    first = 2  # B sütunu
    headers = {first + i: h for i, h in enumerate(HEADERS)}
    if ref_col == "A":
        headers[1] = "ÜYBS REF"
    elif ref_col == "P":
        headers[first + len(HEADERS)] = "ÜYBS REF"
    head_fill, head_font = PatternFill("solid", fgColor="4E8A5B"), Font(bold=True, color="FFFFFF")
    for col, text in headers.items():
        cell = ws.cell(row=1, column=col, value=text)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for i, row in enumerate(template_rows(school, rows), start=2):
        for j, value in enumerate(row["cells"]):
            cell = ws.cell(row=i, column=first + j, value=value)
            if j in (12, 13):
                cell.number_format = "@"  # TC metin olarak kalsın
        if ref_col:
            col = 1 if ref_col == "A" else first + len(HEADERS)
            ref = row["ref"]
            ws.cell(row=i, column=col, value=int(ref) if ref.isdigit() and len(ref) < 15 else ref)
    widths = [10, 10, 12, 10, 11, 11, 11, 14, 8, 8, 12, 12, 8, 16, 18, 10]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    ws.row_dimensions[1].height = 75
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- görünümler ---------------------------------------------------------------------------------

def _week(request):
    try:
        return monday(datetime.date.fromisoformat(request.GET.get("hafta") or request.POST.get("hafta") or ""))
    except ValueError:
        return None


def export(request, pk):
    """ÜYBS şablonunda Excel. Hafta seçiliyse o hafta arşivlenir (arşiv varsa ondan çıkar)."""
    timetable = get_object_or_404(Timetable.objects.select_related("school"), pk=pk)
    week = _week(request)
    ids = [int(x) for x in request.GET.getlist("secim") if x.isdigit()] if request.GET.get("varlik", "sinif") == "sinif" else []
    rows, _ = week_rows(
        timetable, week, ids, substitutions=request.GET.get("vekalet", "1") == "1",
        save=True, user=request.user if request.user.is_authenticated else None,
    )
    response = HttpResponse(
        template_xlsx(timetable.school, rows),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    name = f"uybs-aktarim-{week:%Y-%m-%d}" if week else f"uybs-aktarim-{timetable.pk}"
    response["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    return response


@require_POST
def archive(request, pk):
    """Hafta arşivi: güncel programdan yeniden oluştur ya da sil."""
    timetable = get_object_or_404(Timetable.objects.select_related("school"), pk=pk)
    week = _week(request)
    back = request.POST.get("next") or reverse("timetable:reports", args=[timetable.pk])
    if week is None:
        messages.error(request, "Hafta seçilmedi.")
        return redirect(back)
    if request.POST.get("action") == "delete":
        WeekArchive.objects.filter(school=timetable.school, week_start=week).delete()
        messages.info(request, f"{week:%d.%m.%Y} haftasının arşivi silindi; rapor güncel programdan çıkacak.")
    else:
        archive_week(timetable, week, request.user, substitutions=request.POST.get("vekalet", "1") == "1")
        messages.success(request, f"{week:%d.%m.%Y} haftası “{timetable.name}” sürümünden arşivlendi.")
    return redirect(back)


def _classes_and_teachers(school):
    classes = list(school.classes.all())
    for c in classes:
        c.auto_no = refs.uybs_class_no(c) if not c.uybs_no else ""
    return classes, list(school.teachers.all())


def import_tc(school, text):
    """Yapıştırılan "öğretmen; TC; e-posta" satırları. Öğretmen kısa ad ya da ad soyadla bulunur."""
    by_name = {}
    for t in school.teachers.all():
        for key in (t.full_name, t.short, f"{t.last_name} {t.first_name}"):
            if key:
                by_name.setdefault(key.strip().casefold(), t)
    written, errors = 0, []
    for n, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            continue
        parts = [p.strip() for p in re.split(r"\t|;", line)]
        teacher = by_name.get(parts[0].casefold())
        if teacher is None:
            errors.append(f"{n}. satır: “{parts[0]}” öğretmeni bulunamadı")
            continue
        fields = []
        tc = next((p for p in parts[1:] if TC_RE.fullmatch(p)), None)
        mail = next((p for p in parts[1:] if "@" in p), None)
        if tc:
            teacher.tc_no, fields = tc, fields + ["tc_no"]
        if mail:
            teacher.email, fields = mail[:254], fields + ["email"]
        if not fields:
            errors.append(f"{n}. satır: 11 haneli TC ya da e-posta yok")
            continue
        teacher.save(update_fields=fields)
        written += 1
    return written, errors


def settings_page(request, pk):
    """ÜYBS ayarları: şablon sabitleri, ders kodu önekleri, kısım numaraları, öğretmen TC ve e-postaları."""
    school = get_object_or_404(School, pk=pk)
    settings = settings_of(school)
    prefixes = sorted({prefix_of(s.short) for s in school.subjects.all() if s.short})
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "template":
            for key in ("ksm", "hep1", "hep1b", "hep11"):
                settings[key] = (request.POST.get(key) or "").strip()[:20]
            settings["ref_col"] = request.POST.get("ref_col") if request.POST.get("ref_col") in REF_COLUMNS else "A"
            settings["prefixes"] = {
                p: v for p in prefixes if (v := (request.POST.get(f"prefix-{p}") or "").strip()[:10]) and v != p
            }
            school.uybs_settings = settings
            school.save(update_fields=["uybs_settings"])
            messages.success(request, "Şablon ayarları kaydedildi.")
        elif action == "people":
            bad = []
            with transaction.atomic():
                for c in school.classes.all():
                    value = (request.POST.get(f"class-{c.pk}") or "").strip()[:20]
                    if value != c.uybs_no:
                        c.uybs_no = value
                        c.save(update_fields=["uybs_no"])
                for t in school.teachers.all():
                    tc = (request.POST.get(f"tc-{t.pk}") or "").strip()
                    mail = (request.POST.get(f"mail-{t.pk}") or "").strip()[:254]
                    if tc and not TC_RE.fullmatch(tc):
                        bad.append(t.full_name)
                        tc = t.tc_no
                    if (tc, mail) != (t.tc_no, t.email):
                        t.tc_no, t.email = tc, mail
                        t.save(update_fields=["tc_no", "email"])
            if bad:
                messages.warning(request, "TC 11 haneli olmalı, şunlar kaydedilmedi: " + ", ".join(bad))
            else:
                messages.success(request, "Kısım numaraları ve öğretmen bilgileri kaydedildi.")
        elif action == "import_tc":
            written, errors = import_tc(school, request.POST.get("text", ""))
            messages.success(request, f"{written} öğretmen güncellendi.")
            for e in errors[:10]:
                messages.warning(request, e)
        return redirect("timetable:uybs_settings", pk=school.pk)
    from .views import _nav

    classes, teachers = _classes_and_teachers(school)
    return render(request, "timetable/uybs_settings.html", _nav(
        school, "uybs", settings=settings, ref_columns=REF_COLUMNS, classes=classes, teachers=teachers,
        prefixes=[{"prefix": p, "value": settings["prefixes"].get(p, "")} for p in prefixes],
        missing_tc=sum(1 for t in teachers if not t.tc_no),
        archives=school.week_archives.all()[:30],
    ))
